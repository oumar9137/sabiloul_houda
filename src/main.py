"""Première base du logiciel scolaire Sabiloul Houda (hors ligne)."""

from __future__ import annotations

import sqlite3
import os
import subprocess
import sys
import csv
from datetime import datetime
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

APP_NAME = "Sabiloul Houda — Gestion scolaire"
ROOT = Path(__file__).resolve().parent.parent if not getattr(sys, "frozen", False) else Path(sys.executable).resolve().parent
if getattr(sys, "frozen", False):
    DATA_DIR = Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "SabiloulHouda" / "data"
else:
    DATA_DIR = ROOT / "data"
DB_PATH = DATA_DIR / "sabiloul_houda.sqlite3"
SCHOOL_YEAR = "2026-2027"
MONTHS = ["Octobre", "Novembre", "Décembre", "Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet"]


def asset_path(filename: str) -> Path:
    asset_root = Path(getattr(sys, "_MEIPASS", ROOT))
    return asset_root / "assets" / filename


def connect() -> sqlite3.Connection:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(DB_PATH)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys = ON")
    return db


def initialize_student_matricules(db: sqlite3.Connection) -> None:
    """Migrate existing students once, preserving their registration order."""
    counter = db.execute(
        "SELECT next_value FROM app_counters WHERE name='student_matricule'"
    ).fetchone()
    if counter is None:
        students = db.execute("SELECT id FROM students ORDER BY id").fetchall()
        for row in students:
            db.execute(
                "UPDATE students SET matricule=? WHERE id=?",
                (f"__MATRICULE_MIGRATION_{row['id']}__", row["id"]),
            )
        for number, row in enumerate(students, start=1):
            db.execute(
                "UPDATE students SET matricule=? WHERE id=?",
                (f"S{number:03d}", row["id"]),
            )
        db.execute(
            "INSERT INTO app_counters(name,next_value) VALUES('student_matricule',?)",
            (len(students) + 1,),
        )
        return

    used_numbers = []
    for row in db.execute("SELECT matricule FROM students").fetchall():
        matricule = row["matricule"]
        if matricule.startswith("S") and matricule[1:].isdigit():
            used_numbers.append(int(matricule[1:]))
    next_value = max([counter["next_value"], *(n + 1 for n in used_numbers)])
    db.execute(
        "UPDATE app_counters SET next_value=? WHERE name='student_matricule'",
        (next_value,),
    )


def allocate_student_matricule(db: sqlite3.Connection) -> str:
    row = db.execute(
        "SELECT next_value FROM app_counters WHERE name='student_matricule'"
    ).fetchone()
    if row is None:
        initialize_student_matricules(db)
        row = db.execute(
            "SELECT next_value FROM app_counters WHERE name='student_matricule'"
        ).fetchone()
    number = row["next_value"]
    db.execute(
        "UPDATE app_counters SET next_value=? WHERE name='student_matricule'",
        (number + 1,),
    )
    return f"S{number:03d}"


def initialize_database() -> None:
    with connect() as db:
        db.executescript(
            """
            CREATE TABLE IF NOT EXISTS classes (
                id INTEGER PRIMARY KEY,
                name TEXT NOT NULL UNIQUE,
                active INTEGER NOT NULL DEFAULT 1
            );
            CREATE TABLE IF NOT EXISTS students (
                id INTEGER PRIMARY KEY,
                matricule TEXT NOT NULL UNIQUE,
                surname TEXT NOT NULL,
                given_name TEXT NOT NULL,
                sex TEXT NOT NULL DEFAULT '',
                birth_date TEXT NOT NULL DEFAULT '',
                class_id INTEGER REFERENCES classes(id),
                enrollment_date TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL DEFAULT 'Actif',
                parent_name TEXT NOT NULL DEFAULT '',
                parent_phone TEXT NOT NULL DEFAULT '',
                parent_address TEXT NOT NULL DEFAULT ''
            );
            CREATE TABLE IF NOT EXISTS app_counters (
                name TEXT PRIMARY KEY,
                next_value INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS fees (
                id INTEGER PRIMARY KEY,
                student_id INTEGER NOT NULL REFERENCES students(id),
                school_year TEXT NOT NULL,
                label TEXT NOT NULL,
                amount INTEGER NOT NULL CHECK(amount >= 0),
                UNIQUE(student_id, school_year, label)
            );
            CREATE TABLE IF NOT EXISTS payments (
                id INTEGER PRIMARY KEY,
                receipt_no TEXT NOT NULL UNIQUE,
                student_id INTEGER NOT NULL REFERENCES students(id),
                fee_id INTEGER REFERENCES fees(id),
                paid_at TEXT NOT NULL,
                amount INTEGER NOT NULL CHECK(amount > 0),
                reason TEXT NOT NULL,
                note TEXT NOT NULL DEFAULT ''
            );
            CREATE TABLE IF NOT EXISTS expenses (
                id INTEGER PRIMARY KEY,
                spent_at TEXT NOT NULL,
                category TEXT NOT NULL,
                description TEXT NOT NULL,
                amount INTEGER NOT NULL CHECK(amount > 0),
                payee TEXT NOT NULL DEFAULT '',
                proof TEXT NOT NULL DEFAULT ''
            );
            CREATE TABLE IF NOT EXISTS audit_log (
                id INTEGER PRIMARY KEY,
                happened_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                action TEXT NOT NULL,
                entity TEXT NOT NULL,
                entity_id INTEGER,
                details TEXT NOT NULL DEFAULT ''
            );
            CREATE TABLE IF NOT EXISTS personnel (
                id INTEGER PRIMARY KEY,
                full_name TEXT NOT NULL,
                role TEXT NOT NULL,
                phone TEXT NOT NULL DEFAULT '',
                subject TEXT NOT NULL DEFAULT '',
                classes TEXT NOT NULL DEFAULT '',
                active INTEGER NOT NULL DEFAULT 1
            );
            CREATE TABLE IF NOT EXISTS attendance (
                id INTEGER PRIMARY KEY,
                student_id INTEGER NOT NULL REFERENCES students(id),
                attendance_date TEXT NOT NULL,
                status TEXT NOT NULL CHECK(status IN ('Présent','Absent','Retard')),
                note TEXT NOT NULL DEFAULT '',
                UNIQUE(student_id, attendance_date)
            );
            CREATE TABLE IF NOT EXISTS subjects (
                id INTEGER PRIMARY KEY,
                name TEXT NOT NULL,
                class_id INTEGER NOT NULL REFERENCES classes(id),
                teacher_name TEXT NOT NULL DEFAULT '',
                active INTEGER NOT NULL DEFAULT 1,
                UNIQUE(name, class_id)
            );
            CREATE TABLE IF NOT EXISTS grades (
                id INTEGER PRIMARY KEY,
                student_id INTEGER NOT NULL REFERENCES students(id),
                subject_id INTEGER NOT NULL REFERENCES subjects(id),
                school_year TEXT NOT NULL,
                period TEXT NOT NULL,
                score REAL NOT NULL CHECK(score >= 0),
                max_score REAL NOT NULL CHECK(max_score > 0),
                comment TEXT NOT NULL DEFAULT '',
                UNIQUE(student_id, subject_id, school_year, period)
            );
            """
        )
        initialize_student_matricules(db)
        subject_columns = {row["name"] for row in db.execute("PRAGMA table_info(subjects)")}
        if "active" not in subject_columns:
            db.execute("ALTER TABLE subjects ADD COLUMN active INTEGER NOT NULL DEFAULT 1")
        migrated = db.execute("UPDATE grades SET score=ROUND(score*10.0/max_score,2), max_score=10 WHERE max_score=20").rowcount
        if migrated:
            db.execute(
                "INSERT INTO audit_log(action,entity,details) VALUES('conversion','notes','Anciennes notes sur 20 ramenées sur 10')"
            )
        db.executemany(
            "INSERT OR IGNORE INTO classes(name) VALUES (?)",
            [("Jardin d'enfants",), ("1ère année",), ("2ème année",)],
        )


class SchoolApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title(APP_NAME)
        self.geometry("1080x680")
        self.minsize(850, 560)
        self.configure(bg="#f3f6f4")
        self.logo_image = tk.PhotoImage(file=str(asset_path("logo_sabiloul_houda.png"))).subsample(3, 3)
        self.iconphoto(True, self.logo_image)
        self.style = ttk.Style(self)
        try:
            self.style.theme_use("clam")
        except tk.TclError:
            pass
        self.style.configure("Treeview", rowheight=28, font=("Segoe UI", 10))
        self.style.configure("TFrame", background="#f3f6f4")
        self.style.configure("TLabel", background="#f3f6f4", foreground="#243447")
        self.style.configure("TLabelFrame", background="#ffffff", foreground="#17324d")
        self.style.configure("TLabelFrame.Label", background="#ffffff", foreground="#17324d", font=("Segoe UI", 10, "bold"))
        self.style.configure("Treeview", background="#ffffff", fieldbackground="#ffffff", foreground="#243447")
        self.style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"), background="#17324d", foreground="#ffffff")
        self.style.map("Treeview.Heading", background=[("active", "#24516d")])
        self.style.map("Treeview", background=[("selected", "#167d77")], foreground=[("selected", "#ffffff")])
        self.style.configure("TButton", padding=(8, 4), background="#167d77", foreground="#ffffff")
        self.style.map("TButton", background=[("active", "#11645f"), ("pressed", "#0c504c")], foreground=[("active", "#ffffff")])
        self.style.configure("Danger.TButton", padding=(8, 4), background="#b94e48", foreground="#ffffff")
        self.style.map("Danger.TButton", background=[("active", "#963d39")], foreground=[("active", "#ffffff")])
        self.style.configure("TNotebook", background="#e6eeec", borderwidth=0)
        self.style.configure("TNotebook.Tab", padding=(12, 7), background="#d9e7e4", foreground="#17324d")
        self.style.map("TNotebook.Tab", background=[("selected", "#17324d"), ("active", "#b6d2cc")], foreground=[("selected", "#ffffff")])
        self.style.configure("Accent.TButton", background="#167d77", foreground="#ffffff", padding=(9, 5))
        self.style.map("Accent.TButton", background=[("active", "#11645f")])
        page_colors = {
            "HomePage.TFrame": "#e9eff6", "StudentsPage.TFrame": "#edf5f6",
            "ClassesPage.TFrame": "#eef5ef", "FeesPage.TFrame": "#fff6df",
            "PaymentsPage.TFrame": "#edf2fa", "ExpensesPage.TFrame": "#faeeec",
            "ReportsPage.TFrame": "#eef5f2", "BackupPage.TFrame": "#f1edf7",
            "StaffPage.TFrame": "#edf4f5", "AttendancePage.TFrame": "#f1f5e9",
            "AcademicsPage.TFrame": "#f4f0e7",
        }
        for style_name, color in page_colors.items():
            self.style.configure(style_name, background=color)
        self._build()

    def _build(self) -> None:
        header = ttk.Frame(self, padding=(18, 14))
        header.pack(fill="x")
        header.configure(style="Header.TFrame")
        self.style.configure("Header.TFrame", background="#17324d")
        self.style.configure("HeaderTitle.TLabel", background="#17324d", foreground="#ffffff")
        self.style.configure("HeaderInfo.TLabel", background="#17324d", foreground="#d9e7e4")
        self.style.configure("HeaderLogo.TLabel", background="#17324d")
        ttk.Label(header, image=self.logo_image, style="HeaderLogo.TLabel").pack(side="left", padx=(0, 12))
        ttk.Label(header, text=APP_NAME, font=("Segoe UI", 18, "bold"), style="HeaderTitle.TLabel").pack(side="left")
        ttk.Label(header, text=f"Année scolaire {SCHOOL_YEAR}  •  Fonctionne hors ligne", style="HeaderInfo.TLabel").pack(side="right")

        self.tabs = ttk.Notebook(self)
        self.tabs.pack(fill="both", expand=True, padx=16, pady=(0, 16))
        self.home = ttk.Frame(self.tabs, padding=18, style="HomePage.TFrame")
        self.tabs.add(self.home, text="Accueil")
        self._build_home()
        self.students_tab = ttk.Frame(self.tabs, padding=14, style="StudentsPage.TFrame")
        self.tabs.add(self.students_tab, text="Élèves")
        self._build_students()
        self.classes_tab = ttk.Frame(self.tabs, padding=14, style="ClassesPage.TFrame")
        self.tabs.add(self.classes_tab, text="Classes")
        self._build_classes()
        self.fees_tab = ttk.Frame(self.tabs, padding=14, style="FeesPage.TFrame")
        self.tabs.add(self.fees_tab, text="Frais dus")
        self._build_fees()
        self.payments_tab = ttk.Frame(self.tabs, padding=14, style="PaymentsPage.TFrame")
        self.tabs.add(self.payments_tab, text="Paiements")
        self._build_payments()
        self.expenses_tab = ttk.Frame(self.tabs, padding=14, style="ExpensesPage.TFrame")
        self.tabs.add(self.expenses_tab, text="Dépenses")
        self._build_expenses()
        self.reports_tab = ttk.Frame(self.tabs, padding=14, style="ReportsPage.TFrame")
        self.tabs.add(self.reports_tab, text="Rapports")
        self._build_reports()
        self.backup_tab = ttk.Frame(self.tabs, padding=14, style="BackupPage.TFrame")
        self.tabs.add(self.backup_tab, text="Sauvegarde & journal")
        self._build_backup()
        self.staff_tab = ttk.Frame(self.tabs, padding=14, style="StaffPage.TFrame")
        self.tabs.add(self.staff_tab, text="Personnel")
        self._build_staff()
        self.attendance_tab = ttk.Frame(self.tabs, padding=14, style="AttendancePage.TFrame")
        self.tabs.add(self.attendance_tab, text="Présences")
        self._build_attendance()
        self.academics_tab = ttk.Frame(self.tabs, padding=14, style="AcademicsPage.TFrame")
        self.tabs.add(self.academics_tab, text="Matières & notes")
        self._build_academics()
        self.bulletins_tab = ttk.Frame(self.tabs, padding=14, style="ReportsPage.TFrame")
        self.tabs.add(self.bulletins_tab, text="Bulletins individuels")
        self._build_bulletins()
        self.refresh_all()
        self.refresh_audit_log()

    def _build_staff(self) -> None:
        form = ttk.LabelFrame(self.staff_tab, text="Ajouter un membre du personnel", padding=10)
        form.pack(fill="x", pady=(0, 12))
        self.staff_vars = {key: tk.StringVar() for key in ("name", "role", "phone", "subject", "classes")}
        labels = [("Nom complet *", "name"), ("Fonction *", "role"), ("Téléphone", "phone"), ("Matière", "subject"), ("Classes", "classes")]
        for col, (label, key) in enumerate(labels):
            ttk.Label(form, text=label).grid(row=0, column=col, sticky="w", padx=4)
            ttk.Entry(form, textvariable=self.staff_vars[key], width=23).grid(row=1, column=col, sticky="ew", padx=4)
        ttk.Button(form, text="Enregistrer", command=self.add_staff).grid(row=1, column=5, padx=6)
        self.staff_tree = self._tree(
            self.staff_tab,
            ("name", "role", "phone", "subject", "classes", "state"),
            ("Nom complet", "Fonction", "Téléphone", "Matière", "Classes", "État"),
        )
        buttons = ttk.Frame(self.staff_tab)
        buttons.pack(fill="x", pady=(8, 0))
        ttk.Button(buttons, text="Modifier la fiche", command=self.edit_selected_staff).pack(side="left", padx=(0, 6))
        ttk.Button(buttons, text="Archiver / réactiver", command=self.toggle_selected_staff, style="Danger.TButton").pack(side="left")

    def _build_attendance(self) -> None:
        form = ttk.LabelFrame(self.attendance_tab, text="Enregistrer une présence", padding=10)
        form.pack(fill="x", pady=(0, 12))
        ttk.Label(form, text="Date (AAAA-MM-JJ)").grid(row=0, column=0, sticky="w")
        self.attendance_date = tk.StringVar(value=datetime.now().date().isoformat())
        date_entry = ttk.Entry(form, textvariable=self.attendance_date, width=16)
        date_entry.grid(row=1, column=0, padx=(0, 10), sticky="ew")
        date_entry.bind("<FocusOut>", self.refresh_attendance_students)
        ttk.Label(form, text="Classe").grid(row=0, column=1, sticky="w")
        self.attendance_class = ttk.Combobox(form, state="readonly", width=22)
        self.attendance_class.grid(row=1, column=1, padx=(0, 10), sticky="ew")
        self.attendance_class.bind("<<ComboboxSelected>>", self.refresh_attendance_students)
        ttk.Label(form, text="Élève").grid(row=0, column=2, sticky="w")
        self.attendance_student = ttk.Combobox(form, state="readonly", width=36)
        self.attendance_student.grid(row=1, column=2, padx=(0, 10), sticky="ew")
        ttk.Label(form, text="Statut").grid(row=0, column=3, sticky="w")
        self.attendance_status = ttk.Combobox(form, state="readonly", values=["Présent", "Absent", "Retard"], width=14)
        self.attendance_status.set("Présent")
        self.attendance_status.grid(row=1, column=3, padx=(0, 10), sticky="ew")
        ttk.Label(form, text="Remarque").grid(row=0, column=4, sticky="w")
        self.attendance_note = tk.StringVar()
        ttk.Entry(form, textvariable=self.attendance_note, width=25).grid(row=1, column=4, padx=(0, 10), sticky="ew")
        ttk.Button(form, text="Enregistrer", command=self.add_attendance).grid(row=1, column=5)
        self.attendance_tree = self._tree(
            self.attendance_tab,
            ("date", "matricule", "student", "status", "note"),
            ("Date", "Matricule", "Élève", "Statut", "Remarque"),
        )
        buttons = ttk.Frame(self.attendance_tab)
        buttons.pack(fill="x", pady=(8, 0))
        ttk.Button(buttons, text="Modifier la présence", command=self.edit_selected_attendance).pack(side="left", padx=(0, 6))
        ttk.Button(buttons, text="Supprimer la présence", command=self.delete_selected_attendance, style="Danger.TButton").pack(side="left")

    def _build_academics(self) -> None:
        subject_form = ttk.LabelFrame(self.academics_tab, text="Matières par classe", padding=9)
        subject_form.pack(fill="x", pady=(0, 8))
        ttk.Label(subject_form, text="Classe").grid(row=0, column=0, sticky="w")
        self.subject_class = ttk.Combobox(subject_form, state="readonly", width=22)
        self.subject_class.grid(row=1, column=0, padx=(0, 8), sticky="ew")
        ttk.Label(subject_form, text="Nom de la matière").grid(row=0, column=1, sticky="w")
        self.subject_name = tk.StringVar()
        ttk.Entry(subject_form, textvariable=self.subject_name, width=28).grid(row=1, column=1, padx=(0, 8), sticky="ew")
        ttk.Label(subject_form, text="Enseignant").grid(row=0, column=2, sticky="w")
        self.subject_teacher = tk.StringVar()
        ttk.Entry(subject_form, textvariable=self.subject_teacher, width=24).grid(row=1, column=2, padx=(0, 8), sticky="ew")
        ttk.Button(subject_form, text="Ajouter la matière", command=self.add_subject).grid(row=1, column=3)
        for col in range(3):
            subject_form.columnconfigure(col, weight=1)
        self.subject_tree = self._tree(self.academics_tab, ("class", "subject", "teacher", "state"), ("Classe", "Matière", "Enseignant", "État"))
        self.subject_tree.configure(height=4)
        subject_actions = ttk.Frame(self.academics_tab)
        subject_actions.pack(fill="x", pady=(0, 8))
        ttk.Button(subject_actions, text="Modifier la matière sélectionnée", command=self.edit_selected_subject).pack(side="left", padx=(0, 6))
        ttk.Button(subject_actions, text="Archiver / réactiver la matière", command=self.delete_selected_subject, style="Danger.TButton").pack(side="left")

        grade_form = ttk.LabelFrame(self.academics_tab, text="Enregistrer une note", padding=9)
        grade_form.pack(fill="x", pady=(0, 8))
        ttk.Label(grade_form, text="Élève").grid(row=0, column=0, sticky="w")
        self.grade_student = ttk.Combobox(grade_form, state="readonly", width=31)
        self.grade_student.grid(row=1, column=0, padx=(0, 8), sticky="ew")
        self.grade_student.bind("<<ComboboxSelected>>", self.refresh_grade_subjects)
        ttk.Label(grade_form, text="Matière").grid(row=0, column=1, sticky="w")
        self.grade_subject = ttk.Combobox(grade_form, state="readonly", width=26)
        self.grade_subject.grid(row=1, column=1, padx=(0, 8), sticky="ew")
        ttk.Label(grade_form, text="Trimestre").grid(row=0, column=2, sticky="w")
        self.grade_period = ttk.Combobox(grade_form, state="readonly", values=["Trimestre 1", "Trimestre 2", "Trimestre 3"], width=16)
        self.grade_period.set("Trimestre 1")
        self.grade_period.grid(row=1, column=2, padx=(0, 8), sticky="ew")
        self.grade_period.bind("<<ComboboxSelected>>", lambda _event: self.refresh_grade_list())
        ttk.Label(grade_form, text="Note obtenue").grid(row=0, column=3, sticky="w")
        self.grade_score = tk.StringVar()
        ttk.Entry(grade_form, textvariable=self.grade_score, width=12).grid(row=1, column=3, padx=(0, 8), sticky="ew")
        ttk.Label(grade_form, text="Barème").grid(row=0, column=4, sticky="w")
        self.grade_max = tk.StringVar(value="10")
        ttk.Entry(grade_form, textvariable=self.grade_max, width=10, state="readonly").grid(row=1, column=4, padx=(0, 8), sticky="ew")
        ttk.Label(grade_form, text="Appréciation").grid(row=2, column=0, sticky="w", pady=(6, 0))
        self.grade_comment = tk.StringVar()
        ttk.Entry(grade_form, textvariable=self.grade_comment, width=42).grid(row=3, column=0, columnspan=2, padx=(0, 8), sticky="ew")
        ttk.Label(grade_form, text="Année scolaire").grid(row=2, column=2, sticky="w", pady=(6, 0))
        self.grade_year = tk.StringVar(value=SCHOOL_YEAR)
        grade_year_entry = ttk.Entry(grade_form, textvariable=self.grade_year, width=16)
        grade_year_entry.grid(row=3, column=2, padx=(0, 8), sticky="ew")
        grade_year_entry.bind("<FocusOut>", lambda _event: self.refresh_grade_list())
        ttk.Button(grade_form, text="Enregistrer la note", command=self.add_grade).grid(row=3, column=4, sticky="e")
        for col in range(5):
            grade_form.columnconfigure(col, weight=1)

        self.grade_average = ttk.Label(self.academics_tab, text="Moyenne générale : -", font=("Segoe UI", 11, "bold"))
        self.grade_average.pack(anchor="w", pady=(0, 6))
        self.grade_tree = self._tree(
            self.academics_tab,
            ("subject", "teacher", "score", "max", "normalized", "comment"),
            ("Matière", "Enseignant", "Note", "Sur", "Équivalent /10", "Appréciation"),
        )
        self.grade_tree.configure(height=4)
        grade_actions = ttk.Frame(self.academics_tab)
        grade_actions.pack(fill="x", pady=(6, 0))
        ttk.Button(grade_actions, text="Modifier la note sélectionnée", command=self.edit_selected_grade).pack(side="left", padx=(0, 6))
        ttk.Button(grade_actions, text="Supprimer la note sélectionnée", command=self.delete_selected_grade, style="Danger.TButton").pack(side="left")
        ttk.Button(self.academics_tab, text="Générer le bulletin PDF", command=self.export_bulletin).pack(anchor="w", pady=(8, 0))

    def _build_bulletins(self) -> None:
        form = ttk.LabelFrame(self.bulletins_tab, text="Choisir l'élève et la période", padding=12)
        form.pack(fill="x", pady=(0, 12))
        ttk.Label(form, text="Classe").grid(row=0, column=0, sticky="w", padx=5)
        self.bulletin_class = ttk.Combobox(form, state="readonly", width=24)
        self.bulletin_class.grid(row=1, column=0, padx=5, sticky="ew")
        self.bulletin_class.bind("<<ComboboxSelected>>", self.refresh_bulletin_students)
        ttk.Label(form, text="Trimestre").grid(row=0, column=1, sticky="w", padx=5)
        self.bulletin_period = ttk.Combobox(form, state="readonly", values=["Trimestre 1", "Trimestre 2", "Trimestre 3"], width=20)
        self.bulletin_period.set("Trimestre 1")
        self.bulletin_period.grid(row=1, column=1, padx=5, sticky="ew")
        self.bulletin_period.bind("<<ComboboxSelected>>", self.refresh_bulletin_students)
        ttk.Label(form, text="Année scolaire").grid(row=0, column=2, sticky="w", padx=5)
        self.bulletin_year = tk.StringVar(value=SCHOOL_YEAR)
        year_entry = ttk.Entry(form, textvariable=self.bulletin_year, width=18)
        year_entry.grid(row=1, column=2, padx=5, sticky="ew")
        year_entry.bind("<FocusOut>", self.refresh_bulletin_students)
        form.columnconfigure(0, weight=1)
        form.columnconfigure(1, weight=1)
        self.bulletin_tree = self._tree(self.bulletins_tab, ("matricule", "student", "class", "count", "average"), ("Matricule", "Nom et prénom", "Classe", "Notes", "Moyenne /10"))
        ttk.Button(self.bulletins_tab, text="Générer le bulletin PDF de l'élève sélectionné", command=self.export_selected_bulletin, style="Accent.TButton").pack(anchor="w", pady=10)
        ttk.Label(self.bulletins_tab, text="Chaque bulletin est un document individuel. Sélectionnez un élève, puis choisissez où enregistrer son PDF.").pack(anchor="w")

    def refresh_bulletin_students(self, _event=None) -> None:
        if not hasattr(self, "bulletin_tree"):
            return
        class_name, period, year = self.bulletin_class.get(), self.bulletin_period.get(), self.bulletin_year.get().strip()
        with connect() as db:
            rows = db.execute("""SELECT s.id,s.matricule,s.surname,s.given_name,c.name AS class_name,
                COUNT(g.id) AS grade_count,AVG(g.score*10.0/g.max_score) AS average
                FROM students s LEFT JOIN classes c ON c.id=s.class_id
                LEFT JOIN grades g ON g.student_id=s.id AND g.period=? AND g.school_year=?
                WHERE s.status='Actif' AND c.name=?
                GROUP BY s.id ORDER BY s.surname COLLATE NOCASE,s.given_name COLLATE NOCASE""", (period, year, class_name)).fetchall()
        self.clear(self.bulletin_tree)
        for row in rows:
            average = f"{row['average']:.2f}" if row["average"] is not None else "—"
            self.bulletin_tree.insert("", "end", iid=str(row["id"]), values=(row["matricule"], f"{row['surname']} {row['given_name']}", row["class_name"] or "", row["grade_count"], average))

    def export_selected_bulletin(self) -> None:
        selected = self.bulletin_tree.selection()
        if not selected:
            messagebox.showinfo("Choisir un élève", "Sélectionnez un élève dans la liste.")
            return
        student_id = int(selected[0])
        with connect() as db:
            row = db.execute("SELECT surname,given_name,matricule FROM students WHERE id=?", (student_id,)).fetchone()
        if not row:
            return
        self.grade_student.set(f"{student_id} — {row['surname']} {row['given_name']} ({row['matricule']})")
        self.grade_period.set(self.bulletin_period.get())
        self.grade_year.set(self.bulletin_year.get().strip())
        self.export_bulletin()

    def _build_home(self) -> None:
        ttk.Label(self.home, text="Bienvenue", font=("Segoe UI", 22, "bold")).pack(anchor="w")
        ttk.Label(self.home, text="Gestion locale de l’établissement • Mamaribougou, Mali", foreground="#506070").pack(anchor="w", pady=(4, 20))
        cards = ttk.Frame(self.home)
        cards.pack(fill="x")
        self.metrics: dict[str, ttk.Label] = {}
        for key, label in [("students", "Élèves inscrits"), ("income", "Recettes (FCFA)"), ("expenses", "Dépenses (FCFA)"), ("balance", "Solde caisse (FCFA)")]:
            card = ttk.LabelFrame(cards, text=label, padding=14)
            card.pack(side="left", fill="both", expand=True, padx=(0, 10))
            value = ttk.Label(card, text="0", font=("Segoe UI", 19, "bold"))
            value.pack(anchor="w")
            self.metrics[key] = value
        ttk.Label(self.home, text="Les données restent sur cet ordinateur. Configurez les frais par élève avant d'enregistrer les règlements.", wraplength=800).pack(anchor="w", pady=28)

    def _build_students(self) -> None:
        form = ttk.LabelFrame(self.students_tab, text="Ajouter un élève", padding=10)
        form.pack(fill="x", pady=(0, 12))
        self.student_vars = {k: tk.StringVar() for k in ("surname", "given_name", "sex", "birth_date", "class_name", "enrollment_date", "parent_name", "parent_phone")}
        ttk.Label(form, text="Le matricule S001, S002… est attribué automatiquement à l'inscription.").grid(row=0, column=0, columnspan=4, sticky="w", padx=4, pady=(0, 5))
        fields = [("Nom *", "surname"), ("Prénom *", "given_name"), ("Sexe", "sex"), ("Naissance (AAAA-MM-JJ)", "birth_date"), ("Date inscription", "enrollment_date"), ("Responsable", "parent_name"), ("Téléphone", "parent_phone")]
        for idx, (label, key) in enumerate(fields):
            row, col = divmod(idx, 4)
            row += 1
            box = ttk.Frame(form)
            box.grid(row=row, column=col, sticky="ew", padx=4, pady=3)
            ttk.Label(box, text=label).pack(anchor="w")
            ttk.Entry(box, textvariable=self.student_vars[key], width=22).pack(fill="x")
        ttk.Label(form, text="Classe").grid(row=3, column=0, sticky="w", padx=4)
        self.student_class = ttk.Combobox(form, state="readonly", width=20)
        self.student_class.grid(row=4, column=0, sticky="ew", padx=4, pady=3)
        ttk.Button(form, text="Enregistrer l'élève", command=self.add_student, style="Accent.TButton").grid(row=4, column=1, sticky="ew", padx=4)
        ttk.Button(form, text="Modifier l'élève", command=self.edit_selected_student).grid(row=4, column=2, sticky="ew", padx=4)
        ttk.Button(form, text="Supprimer l'élève", command=self.permanently_delete_selected_student, style="Danger.TButton").grid(row=4, column=3, sticky="ew", padx=4)
        for col in range(4):
            form.columnconfigure(col, weight=1)
        filter_row = ttk.Frame(self.students_tab)
        filter_row.pack(fill="x", pady=(0, 8))
        ttk.Label(filter_row, text="Afficher les élèves de la classe :").pack(side="left", padx=(0, 8))
        self.student_filter_class = ttk.Combobox(filter_row, state="readonly", width=28)
        self.student_filter_class.pack(side="left")
        self.student_filter_class.bind("<<ComboboxSelected>>", lambda _event: self.refresh_student_list())
        self.student_tree = self._tree(self.students_tab, ("matricule", "nom", "prenom", "classe", "telephone", "statut"), ("Matricule", "Nom", "Prénom", "Classe", "Téléphone", "Statut"))
        actions = ttk.Frame(self.students_tab)
        actions.pack(fill="x", pady=(8, 0))
        ttk.Button(actions, text="Archiver / réactiver l'élève", command=self.delete_selected_student).pack(side="left", padx=(0, 6))

    def _build_classes(self) -> None:
        row = ttk.Frame(self.classes_tab)
        row.pack(fill="x", pady=(0, 12))
        ttk.Label(row, text="Nouvelle classe").pack(side="left")
        self.class_name = tk.StringVar()
        ttk.Entry(row, textvariable=self.class_name, width=30).pack(side="left", padx=8)
        ttk.Button(row, text="Ajouter", command=self.add_class).pack(side="left")
        self.class_tree = self._tree(self.classes_tab, ("name", "count", "state"), ("Classe", "Nombre d'élèves", "État"))
        buttons = ttk.Frame(self.classes_tab)
        buttons.pack(fill="x", pady=(8, 0))
        ttk.Button(buttons, text="Modifier la classe", command=self.edit_selected_class).pack(side="left", padx=(0, 6))
        ttk.Button(buttons, text="Archiver / réactiver", command=self.toggle_selected_class, style="Danger.TButton").pack(side="left")

    def _build_fees(self) -> None:
        form = ttk.LabelFrame(self.fees_tab, text="Ajouter un frais dû à un élève", padding=10)
        form.pack(fill="x", pady=(0, 12))
        ttk.Label(form, text="Élève").grid(row=0, column=0, sticky="w")
        self.fee_student = ttk.Combobox(form, state="readonly", width=38)
        self.fee_student.grid(row=1, column=0, padx=(0, 10), sticky="ew")
        ttk.Label(form, text="Année scolaire").grid(row=0, column=1, sticky="w")
        self.fee_year = tk.StringVar(value=SCHOOL_YEAR)
        ttk.Entry(form, textvariable=self.fee_year, width=14).grid(row=1, column=1, padx=(0, 10), sticky="ew")
        ttk.Label(form, text="Frais / mensualité").grid(row=0, column=2, sticky="w")
        self.fee_label = ttk.Combobox(form, state="normal", values=["Inscription", *MONTHS], width=18)
        self.fee_label.grid(row=1, column=2, padx=(0, 10), sticky="ew")
        ttk.Label(form, text="Montant dû (FCFA)").grid(row=0, column=3, sticky="w")
        self.fee_amount = tk.StringVar()
        ttk.Entry(form, textvariable=self.fee_amount, width=16).grid(row=1, column=3, padx=(0, 10), sticky="ew")
        ttk.Button(form, text="Ajouter le frais", command=self.add_fee).grid(row=1, column=4)
        for col in range(4):
            form.columnconfigure(col, weight=1)
        self.fee_tree = self._tree(
            self.fees_tab,
            ("student", "year", "label", "due", "paid", "remaining"),
            ("Élève", "Année", "Frais / mois", "Montant dû", "Payé", "Reste à payer (FCFA)"),
        )
        buttons = ttk.Frame(self.fees_tab)
        buttons.pack(fill="x", pady=(8, 0))
        ttk.Button(buttons, text="Modifier les frais", command=self.edit_selected_fee).pack(side="left", padx=(0, 6))
        ttk.Button(buttons, text="Supprimer les frais", command=self.delete_selected_fee, style="Danger.TButton").pack(side="left")

    def _build_payments(self) -> None:
        form = ttk.LabelFrame(self.payments_tab, text="Enregistrer un paiement", padding=10)
        form.pack(fill="x", pady=(0, 12))
        ttk.Label(form, text="Frais à régler").grid(row=0, column=0, sticky="w")
        self.payment_fee = ttk.Combobox(form, state="readonly", width=54)
        self.payment_fee.grid(row=1, column=0, padx=(0, 12), sticky="ew")
        self.payment_fee.bind("<<ComboboxSelected>>", self.update_payment_balance)
        self.payment_amount = tk.StringVar()
        ttk.Label(form, text="Montant versé (FCFA)").grid(row=0, column=1, sticky="w")
        ttk.Entry(form, textvariable=self.payment_amount, width=18).grid(row=1, column=1, padx=(0, 12), sticky="ew")
        self.payment_balance = ttk.Label(form, text="Reste à payer : -")
        self.payment_balance.grid(row=0, column=2, sticky="w", padx=(0, 12))
        ttk.Button(form, text="Enregistrer", command=self.add_payment).grid(row=1, column=2, sticky="w")
        ttk.Button(self.payments_tab, text="Créer le reçu PDF du paiement sélectionné", command=self.export_selected_receipt).pack(anchor="w", pady=(0, 10))
        self.payment_tree = self._tree(self.payments_tab, ("receipt", "date", "student", "reason", "amount"), ("Reçu", "Date", "Élève", "Motif", "Montant (FCFA)"))
        buttons = ttk.Frame(self.payments_tab)
        buttons.pack(fill="x", pady=(8, 0))
        ttk.Button(buttons, text="Modifier le paiement", command=self.edit_selected_payment).pack(side="left", padx=(0, 6))
        ttk.Button(buttons, text="Supprimer le paiement", command=self.delete_selected_payment, style="Danger.TButton").pack(side="left")

    def _build_expenses(self) -> None:
        form = ttk.LabelFrame(self.expenses_tab, text="Enregistrer une dépense", padding=10)
        form.pack(fill="x", pady=(0, 12))
        self.expense_vars = {k: tk.StringVar() for k in ("date", "category", "description", "amount", "payee")}
        labels = [("Date (AAAA-MM-JJ)", "date"), ("Catégorie", "category"), ("Description", "description"), ("Montant (FCFA)", "amount"), ("Bénéficiaire", "payee")]
        for i, (label, key) in enumerate(labels):
            ttk.Label(form, text=label).grid(row=0, column=i, sticky="w", padx=3)
            ttk.Entry(form, textvariable=self.expense_vars[key], width=22).grid(row=1, column=i, sticky="ew", padx=3)
        ttk.Button(form, text="Enregistrer", command=self.add_expense).grid(row=1, column=5, padx=6)
        self.expense_tree = self._tree(self.expenses_tab, ("date", "category", "description", "payee", "amount"), ("Date", "Catégorie", "Description", "Bénéficiaire", "Montant (FCFA)"))
        expense_actions = ttk.Frame(self.expenses_tab)
        expense_actions.pack(fill="x", pady=(8, 0))
        ttk.Button(expense_actions, text="Modifier la dépense sélectionnée", command=self.edit_selected_expense).pack(side="left", padx=(0, 6))
        ttk.Button(expense_actions, text="Supprimer la dépense sélectionnée", command=self.delete_selected_expense, style="Danger.TButton").pack(side="left")

    def _build_reports(self) -> None:
        toolbar = ttk.Frame(self.reports_tab)
        toolbar.pack(fill="x", pady=(0, 10))
        ttk.Label(toolbar, text="Rapport").pack(side="left")
        self.report_choice = ttk.Combobox(
            toolbar,
            state="readonly",
            width=28,
            values=["Liste des élèves", "Élèves par classe (A à Z)", "Matières par classe", "Frais impayés", "Paiements", "Dépenses", "Présences", "Notes"],
        )
        self.report_choice.set("Liste des élèves")
        self.report_choice.pack(side="left", padx=8)
        ttk.Label(toolbar, text="Classe").pack(side="left", padx=(4, 0))
        self.report_class_filter = ttk.Combobox(toolbar, state="readonly", width=22)
        self.report_class_filter.pack(side="left", padx=8)
        self.report_choice.bind("<<ComboboxSelected>>", lambda _event: self.update_report_class_filter())
        ttk.Button(toolbar, text="Afficher", command=self.load_report).pack(side="left", padx=(0, 16))
        ttk.Button(toolbar, text="Exporter CSV", command=self.export_report_csv).pack(side="left", padx=4)
        ttk.Button(toolbar, text="Exporter Excel", command=self.export_report_xlsx).pack(side="left", padx=4)
        self.report_status = ttk.Label(self.reports_tab, text="Choisissez un rapport puis cliquez sur Afficher.", foreground="#506070")
        self.report_status.pack(anchor="w", pady=(0, 6))
        frame = ttk.Frame(self.reports_tab)
        frame.pack(fill="both", expand=True)
        self.report_tree = ttk.Treeview(frame, show="headings")
        vertical = ttk.Scrollbar(frame, orient="vertical", command=self.report_tree.yview)
        horizontal = ttk.Scrollbar(frame, orient="horizontal", command=self.report_tree.xview)
        self.report_tree.configure(yscrollcommand=vertical.set, xscrollcommand=horizontal.set)
        self.report_tree.grid(row=0, column=0, sticky="nsew")
        vertical.grid(row=0, column=1, sticky="ns")
        horizontal.grid(row=1, column=0, sticky="ew")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        self.report_headers: list[str] = []
        self.report_rows: list[tuple] = []
        self.report_title = ""

    def update_report_class_filter(self) -> None:
        if self.report_choice.get() == "Dépenses":
            self.report_class_filter.set("Toutes les classes")
            self.report_class_filter.configure(state="disabled")
        else:
            self.report_class_filter.configure(state="readonly")

    def _build_backup(self) -> None:
        actions = ttk.LabelFrame(self.backup_tab, text="Sauvegarde et restauration", padding=12)
        actions.pack(fill="x", pady=(0, 14))
        ttk.Button(actions, text="Créer une sauvegarde", command=self.create_backup).pack(side="left", padx=(0, 8))
        ttk.Button(actions, text="Restaurer une sauvegarde", command=self.restore_backup).pack(side="left", padx=8)
        ttk.Label(
            actions,
            text="La restauration remplace les données actuelles. Une copie de sécurité sera créée automatiquement.",
            wraplength=620,
        ).pack(side="left", padx=18)
        history_header = ttk.Frame(self.backup_tab)
        history_header.pack(fill="x", pady=(0, 8))
        ttk.Label(history_header, text="Journal des opérations", font=("Segoe UI", 12, "bold")).pack(side="left")
        ttk.Button(history_header, text="Actualiser", command=self.refresh_audit_log).pack(side="right")
        self.audit_tree = self._tree(
            self.backup_tab,
            ("date", "action", "entity", "entity_id", "details"),
            ("Date et heure", "Action", "Élément", "N°", "Détails"),
        )

    @staticmethod
    def _tree(parent: ttk.Frame, keys: tuple[str, ...], headings: tuple[str, ...]) -> ttk.Treeview:
        frame = ttk.Frame(parent)
        frame.pack(fill="both", expand=True)
        tree = ttk.Treeview(frame, columns=keys, show="headings")
        scroll = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=scroll.set)
        for key, heading in zip(keys, headings):
            tree.heading(key, text=heading)
            tree.column(key, width=140, anchor="w")
        tree.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")
        return tree

    def _edit_dialog(self, title: str, fields: list[dict], on_save) -> None:
        dialog = tk.Toplevel(self)
        dialog.title(title)
        dialog.transient(self)
        dialog.grab_set()
        variables = {}
        for index, field in enumerate(fields):
            ttk.Label(dialog, text=field["label"]).grid(row=index, column=0, padx=10, pady=5, sticky="w")
            variable = tk.StringVar(value=str(field.get("value", "")))
            variables[field["key"]] = variable
            if field.get("choices") is not None:
                control = ttk.Combobox(dialog, textvariable=variable, values=field["choices"], state="readonly", width=32)
            else:
                control = ttk.Entry(dialog, textvariable=variable, width=35)
            control.grid(row=index, column=1, padx=10, pady=5, sticky="ew")

        def save():
            if on_save({key: value.get().strip() for key, value in variables.items()}, dialog):
                dialog.destroy()

        ttk.Button(dialog, text="Enregistrer les changements", command=save, style="Accent.TButton").grid(row=len(fields), column=0, columnspan=2, padx=10, pady=10)

    def add_student(self) -> None:
        data = {k: v.get().strip() for k, v in self.student_vars.items()}
        if not data["surname"] or not data["given_name"]:
            messagebox.showwarning("Champs requis", "Le nom et le prénom sont obligatoires.")
            return
        try:
            with connect() as db:
                matricule = allocate_student_matricule(db)
                class_row = db.execute("SELECT id FROM classes WHERE name=?", (self.student_class.get(),)).fetchone()
                cur = db.execute("""INSERT INTO students(matricule,surname,given_name,sex,birth_date,class_id,enrollment_date,parent_name,parent_phone)
                    VALUES(?,?,?,?,?,?,?,?,?)""", (matricule, data["surname"], data["given_name"], data["sex"], data["birth_date"], class_row["id"] if class_row else None, data["enrollment_date"], data["parent_name"], data["parent_phone"]))
                db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('création','élève',?,?)", (cur.lastrowid, matricule))
        except sqlite3.IntegrityError:
            messagebox.showerror("Inscription impossible", "L'élève n'a pas pu être enregistré. Vérifiez les informations saisies.")
            return
        for var in self.student_vars.values():
            var.set("")
        self.refresh_all()

    def edit_selected_student(self) -> None:
        selected = self.student_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez d'abord un élève dans la liste.")
            return
        student_id = int(selected[0])
        with connect() as db:
            row = db.execute("SELECT * FROM students WHERE id=?", (student_id,)).fetchone()
            classes = [r["name"] for r in db.execute("SELECT name FROM classes WHERE active=1 ORDER BY name").fetchall()]
            class_row = db.execute("SELECT name FROM classes WHERE id=?", (row["class_id"],)).fetchone() if row["class_id"] else None
        if not row:
            return
        dialog = tk.Toplevel(self)
        dialog.title("Modifier l'élève")
        dialog.transient(self)
        dialog.grab_set()
        fields = [("Matricule (automatique)", "matricule"), ("Nom", "surname"), ("Prénom", "given_name"), ("Sexe", "sex"), ("Naissance", "birth_date"), ("Inscription", "enrollment_date"), ("Responsable", "parent_name"), ("Téléphone", "parent_phone")]
        variables = {}
        for idx, (label, key) in enumerate(fields):
            ttk.Label(dialog, text=label).grid(row=idx, column=0, sticky="w", padx=10, pady=4)
            var = tk.StringVar(value=row[key] or "")
            variables[key] = var
            ttk.Entry(dialog, textvariable=var, width=34, state="readonly" if key == "matricule" else "normal").grid(row=idx, column=1, padx=10, pady=4)
        ttk.Label(dialog, text="Classe").grid(row=len(fields), column=0, sticky="w", padx=10, pady=4)
        class_var = tk.StringVar(value=class_row["name"] if class_row else "")
        ttk.Combobox(dialog, textvariable=class_var, values=classes, state="readonly", width=31).grid(row=len(fields), column=1, padx=10, pady=4)

        def save():
            data = {key: var.get().strip() for key, var in variables.items()}
            if not data["surname"] or not data["given_name"]:
                messagebox.showwarning("Champs requis", "Le nom et le prénom sont obligatoires.", parent=dialog)
                return
            with connect() as db:
                cls = db.execute("SELECT id FROM classes WHERE name=?", (class_var.get(),)).fetchone()
                try:
                    db.execute("""UPDATE students SET surname=?,given_name=?,sex=?,birth_date=?,class_id=?,enrollment_date=?,parent_name=?,parent_phone=? WHERE id=?""",
                        (data["surname"], data["given_name"], data["sex"], data["birth_date"], cls["id"] if cls else None, data["enrollment_date"], data["parent_name"], data["parent_phone"], student_id))
                except sqlite3.IntegrityError:
                    messagebox.showerror("Modification impossible", "Vérifiez les informations saisies.", parent=dialog)
                    return
                db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('modification','élève',?,?)", (student_id, row["matricule"]))
            dialog.destroy()
            self.refresh_all()

        ttk.Button(dialog, text="Enregistrer les changements", command=save, style="Accent.TButton").grid(row=len(fields)+1, column=0, columnspan=2, pady=10)

    def delete_selected_student(self) -> None:
        selected = self.student_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez d'abord un élève dans la liste.")
            return
        student_id = int(selected[0])
        with connect() as db:
            record = db.execute("SELECT matricule,surname,given_name,status FROM students WHERE id=?", (student_id,)).fetchone()
        if not record:
            return
        action = "archiver" if record["status"] != "Archivé" else "réactiver"
        if not messagebox.askyesno("Confirmer", f"Voulez-vous {action} la fiche de {record['surname']} {record['given_name']} ?", icon="warning"):
            return
        new_status = "Archivé" if record["status"] != "Archivé" else "Actif"
        with connect() as db:
            db.execute("UPDATE students SET status=? WHERE id=?", (new_status, student_id))
            db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES(?,?,?,?)", ("archivage" if new_status == "Archivé" else "réactivation", "élève", student_id, record["matricule"]))
        self.refresh_all()

    def permanently_delete_selected_student(self) -> None:
        selected = self.student_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez d'abord un élève dans la liste.")
            return
        student_id = int(selected[0])
        with connect() as db:
            record = db.execute("SELECT matricule,surname,given_name FROM students WHERE id=?", (student_id,)).fetchone()
            if not record:
                return
            counts = {
                "notes": db.execute("SELECT COUNT(*) FROM grades WHERE student_id=?", (student_id,)).fetchone()[0],
                "présences": db.execute("SELECT COUNT(*) FROM attendance WHERE student_id=?", (student_id,)).fetchone()[0],
                "frais": db.execute("SELECT COUNT(*) FROM fees WHERE student_id=?", (student_id,)).fetchone()[0],
                "paiements": db.execute("SELECT COUNT(*) FROM payments WHERE student_id=?", (student_id,)).fetchone()[0],
            }
        details = ", ".join(f"{n} {label}" for label, n in counts.items() if n)
        related = details or "aucune note, présence, frais ou paiement"
        confirmed = messagebox.askyesno(
            "Suppression définitive",
            f"Supprimer définitivement {record['surname']} {record['given_name']} ({record['matricule']}) ?\n\n"
            f"Cette action effacera sa fiche et {related}. Elle ne peut pas être annulée.\n\nContinuer ?",
            icon="warning",
        )
        if not confirmed:
            return
        with connect() as db:
            db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('suppression définitive','élève',?,?)", (student_id, record["matricule"]))
            db.execute("DELETE FROM grades WHERE student_id=?", (student_id,))
            db.execute("DELETE FROM attendance WHERE student_id=?", (student_id,))
            db.execute("DELETE FROM payments WHERE student_id=?", (student_id,))
            db.execute("DELETE FROM fees WHERE student_id=?", (student_id,))
            db.execute("DELETE FROM students WHERE id=?", (student_id,))
        self.refresh_all()

    def edit_selected_grade(self) -> None:
        selected = self.grade_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez une note dans la liste.")
            return
        grade_id = int(selected[0])
        with connect() as db:
            row = db.execute("SELECT score,max_score,comment FROM grades WHERE id=?", (grade_id,)).fetchone()
        if not row:
            return
        dialog = tk.Toplevel(self)
        dialog.title("Modifier la note")
        dialog.transient(self)
        dialog.grab_set()
        score_var, comment_var = tk.StringVar(value=str(row["score"])), tk.StringVar(value=row["comment"])
        ttk.Label(dialog, text=f"Note sur {row['max_score']}").grid(row=0, column=0, padx=10, pady=6)
        ttk.Entry(dialog, textvariable=score_var, width=14).grid(row=0, column=1, padx=10, pady=6)
        ttk.Label(dialog, text="Appréciation").grid(row=1, column=0, padx=10, pady=6)
        ttk.Entry(dialog, textvariable=comment_var, width=32).grid(row=1, column=1, padx=10, pady=6)

        def save():
            try:
                score = float(score_var.get().replace(",", "."))
                if score < 0 or score > row["max_score"]:
                    raise ValueError
            except ValueError:
                messagebox.showwarning("Note invalide", f"La note doit être comprise entre 0 et {row['max_score']}.", parent=dialog)
                return
            with connect() as db:
                db.execute("UPDATE grades SET score=?,comment=? WHERE id=?", (score, comment_var.get().strip(), grade_id))
                db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('modification','note',?,?)", (grade_id, f"{score}/{row['max_score']}"))
            dialog.destroy()
            self.refresh_grade_list()

        ttk.Button(dialog, text="Enregistrer", command=save, style="Accent.TButton").grid(row=2, column=0, columnspan=2, pady=8)

    def delete_selected_grade(self) -> None:
        selected = self.grade_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez une note dans la liste.")
            return
        grade_id = int(selected[0])
        if not messagebox.askyesno("Confirmer la suppression", "Supprimer cette note ?", icon="warning"):
            return
        with connect() as db:
            db.execute("DELETE FROM grades WHERE id=?", (grade_id,))
            db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('suppression','note',?,?)", (grade_id, "Note supprimée"))
        self.refresh_grade_list()

    def add_class(self) -> None:
        name = self.class_name.get().strip()
        if not name:
            return
        try:
            with connect() as db:
                cur = db.execute("INSERT INTO classes(name) VALUES(?)", (name,))
                db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('création','classe',?,?)", (cur.lastrowid, name))
        except sqlite3.IntegrityError:
            messagebox.showwarning("Classe existante", "Cette classe existe déjà.")
        self.class_name.set("")
        self.refresh_all()

    def edit_selected_class(self) -> None:
        selected = self.class_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez une classe dans la liste.")
            return
        class_id = int(selected[0])
        with connect() as db:
            row = db.execute("SELECT name,active FROM classes WHERE id=?", (class_id,)).fetchone()
        if not row:
            return

        def save(values, dialog):
            name = values["name"].strip()
            active = 1 if values["state"] == "Active" else 0
            if not name:
                messagebox.showwarning("Nom requis", "Indiquez le nom de la classe.", parent=dialog)
                return False
            try:
                with connect() as db:
                    db.execute("UPDATE classes SET name=?,active=? WHERE id=?", (name, active, class_id))
                    db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('modification','classe',?,?)", (class_id, f"{name}; {'active' if active else 'archivée'}"))
            except sqlite3.IntegrityError:
                messagebox.showerror("Nom déjà utilisé", "Une classe porte déjà ce nom.", parent=dialog)
                return False
            self.refresh_all()
            return True

        self._edit_dialog("Modifier la classe", [
            {"label": "Nom de la classe", "key": "name", "value": row["name"]},
            {"label": "État", "key": "state", "value": "Active" if row["active"] else "Archivée", "choices": ["Active", "Archivée"]},
        ], save)

    def toggle_selected_class(self) -> None:
        selected = self.class_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez une classe dans la liste.")
            return
        class_id = int(selected[0])
        with connect() as db:
            row = db.execute("SELECT name,active FROM classes WHERE id=?", (class_id,)).fetchone()
            if not row:
                return
            action = "archiver" if row["active"] else "réactiver"
            if not messagebox.askyesno("Confirmer", f"Voulez-vous {action} la classe « {row['name']} » ?", icon="warning"):
                return
            active = 0 if row["active"] else 1
            db.execute("UPDATE classes SET active=? WHERE id=?", (active, class_id))
            db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES(?,?,?,?)", ("archivage" if not active else "réactivation", "classe", class_id, row["name"]))
        self.refresh_all()

    def add_staff(self) -> None:
        data = {key: value.get().strip() for key, value in self.staff_vars.items()}
        if not data["name"] or not data["role"]:
            messagebox.showwarning("Champs requis", "Le nom complet et la fonction sont obligatoires.")
            return
        with connect() as db:
            cur = db.execute(
                "INSERT INTO personnel(full_name,role,phone,subject,classes) VALUES(?,?,?,?,?)",
                (data["name"], data["role"], data["phone"], data["subject"], data["classes"]),
            )
            db.execute(
                "INSERT INTO audit_log(action,entity,entity_id,details) VALUES('création','personnel',?,?)",
                (cur.lastrowid, f"{data['name']} - {data['role']}"),
            )
        for value in self.staff_vars.values():
            value.set("")
        self.refresh_staff()

    def refresh_staff(self) -> None:
        with connect() as db:
            rows = db.execute("SELECT id,full_name,role,phone,subject,classes,active FROM personnel ORDER BY full_name COLLATE NOCASE").fetchall()
        self.clear(self.staff_tree)
        for row in rows:
            self.staff_tree.insert("", "end", iid=str(row["id"]), values=(row["full_name"], row["role"], row["phone"], row["subject"], row["classes"], "Actif" if row["active"] else "Archivé"))

    def edit_selected_staff(self) -> None:
        selected = self.staff_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez une fiche du personnel.")
            return
        staff_id = int(selected[0])
        with connect() as db:
            row = db.execute("SELECT full_name,role,phone,subject,classes,active FROM personnel WHERE id=?", (staff_id,)).fetchone()
        if not row:
            return

        def save(values, dialog):
            if not values["full_name"] or not values["role"]:
                messagebox.showwarning("Champs requis", "Le nom et la fonction sont obligatoires.", parent=dialog)
                return False
            active = 1 if values["state"] == "Actif" else 0
            with connect() as db:
                db.execute("UPDATE personnel SET full_name=?,role=?,phone=?,subject=?,classes=?,active=? WHERE id=?", (values["full_name"], values["role"], values["phone"], values["subject"], values["classes"], active, staff_id))
                db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('modification','personnel',?,?)", (staff_id, values["full_name"]))
            self.refresh_staff()
            return True

        fields = [
            {"label": "Nom complet", "key": "full_name", "value": row["full_name"]},
            {"label": "Fonction", "key": "role", "value": row["role"]},
            {"label": "Téléphone", "key": "phone", "value": row["phone"]},
            {"label": "Matière", "key": "subject", "value": row["subject"]},
            {"label": "Classes", "key": "classes", "value": row["classes"]},
            {"label": "État", "key": "state", "value": "Actif" if row["active"] else "Archivé", "choices": ["Actif", "Archivé"]},
        ]
        self._edit_dialog("Modifier le personnel", fields, save)

    def toggle_selected_staff(self) -> None:
        selected = self.staff_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez une fiche du personnel.")
            return
        staff_id = int(selected[0])
        with connect() as db:
            row = db.execute("SELECT full_name,active FROM personnel WHERE id=?", (staff_id,)).fetchone()
            if not row:
                return
            action = "archiver" if row["active"] else "réactiver"
            if not messagebox.askyesno("Confirmer", f"Voulez-vous {action} la fiche de {row['full_name']} ?", icon="warning"):
                return
            active = 0 if row["active"] else 1
            db.execute("UPDATE personnel SET active=? WHERE id=?", (active, staff_id))
            db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES(?,?,?,?)", ("archivage" if not active else "réactivation", "personnel", staff_id, row["full_name"]))
        self.refresh_staff()

    def refresh_attendance_students(self, _event=None) -> None:
        class_name = self.attendance_class.get()
        with connect() as db:
            rows = db.execute("""SELECT s.id,s.matricule,s.surname,s.given_name FROM students s
                JOIN classes c ON c.id=s.class_id WHERE c.name=? AND c.active=1 AND s.status='Actif' ORDER BY s.surname,s.given_name""", (class_name,)).fetchall()
        self.attendance_student["values"] = [f"{row['id']} — {row['surname']} {row['given_name']} ({row['matricule']})" for row in rows]
        if self.attendance_student.get() not in self.attendance_student["values"]:
            self.attendance_student.set(self.attendance_student["values"][0] if self.attendance_student["values"] else "")
        self.refresh_attendance_table()

    def add_subject(self) -> None:
        class_name = self.subject_class.get()
        name = self.subject_name.get().strip()
        teacher = self.subject_teacher.get().strip()
        if not class_name or not name:
            messagebox.showwarning("Matière incomplète", "Choisissez une classe et indiquez le nom de la matière.")
            return
        with connect() as db:
            class_row = db.execute("SELECT id FROM classes WHERE name=?", (class_name,)).fetchone()
            try:
                cur = db.execute("INSERT INTO subjects(name,class_id,teacher_name) VALUES(?,?,?)", (name, class_row["id"], teacher))
            except sqlite3.IntegrityError:
                messagebox.showwarning("Matière déjà enregistrée", "Cette matière existe déjà dans cette classe.")
                return
            db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('création','matière',?,?)", (cur.lastrowid, f"{name}; {class_name}"))
        self.subject_name.set("")
        self.subject_teacher.set("")
        self.refresh_academics()

    def refresh_academics(self) -> None:
        with connect() as db:
            classes = [row["name"] for row in db.execute("SELECT name FROM classes WHERE active=1 ORDER BY name").fetchall()]
            students = db.execute("SELECT id,matricule,surname,given_name FROM students WHERE status='Actif' ORDER BY surname,given_name").fetchall()
            subjects = db.execute("""SELECT sub.id,c.name AS class_name,sub.name,sub.teacher_name,sub.active
                FROM subjects sub JOIN classes c ON c.id=sub.class_id
                ORDER BY c.name COLLATE NOCASE,sub.name COLLATE NOCASE""").fetchall()
        self.clear(self.subject_tree)
        for subject in subjects:
            self.subject_tree.insert("", "end", iid=str(subject["id"]), values=(subject["class_name"], subject["name"], subject["teacher_name"], "Active" if subject["active"] else "Archivée"))
        self.subject_class["values"] = classes
        if self.subject_class.get() not in classes:
            self.subject_class.set(classes[0] if classes else "")
        student_choices = [f"{row['id']} — {row['surname']} {row['given_name']} ({row['matricule']})" for row in students]
        self.grade_student["values"] = student_choices
        if self.grade_student.get() not in student_choices:
            self.grade_student.set(student_choices[0] if student_choices else "")
        self.refresh_grade_subjects()

    def edit_selected_subject(self) -> None:
        selected = self.subject_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez une matière dans la liste.")
            return
        subject_id = int(selected[0])
        with connect() as db:
            row = db.execute("SELECT sub.name,sub.teacher_name,sub.active,c.name AS class_name FROM subjects sub JOIN classes c ON c.id=sub.class_id WHERE sub.id=?", (subject_id,)).fetchone()
            classes = [r["name"] for r in db.execute("SELECT name FROM classes WHERE active=1 ORDER BY name").fetchall()]
        if not row:
            return
        dialog = tk.Toplevel(self)
        dialog.title("Modifier la matière")
        dialog.transient(self)
        dialog.grab_set()
        name_var, teacher_var, class_var = tk.StringVar(value=row["name"]), tk.StringVar(value=row["teacher_name"]), tk.StringVar(value=row["class_name"])
        for idx, (label, var) in enumerate((("Matière", name_var), ("Enseignant", teacher_var))):
            ttk.Label(dialog, text=label).grid(row=idx, column=0, padx=10, pady=5, sticky="w")
            ttk.Entry(dialog, textvariable=var, width=30).grid(row=idx, column=1, padx=10, pady=5)
        ttk.Label(dialog, text="Classe").grid(row=2, column=0, padx=10, pady=5, sticky="w")
        ttk.Combobox(dialog, textvariable=class_var, values=classes, state="readonly", width=27).grid(row=2, column=1, padx=10, pady=5)
        ttk.Label(dialog, text="État").grid(row=3, column=0, padx=10, pady=5, sticky="w")
        active_var = tk.StringVar(value="Active" if row["active"] else "Archivée")
        ttk.Combobox(dialog, textvariable=active_var, values=["Active", "Archivée"], state="readonly", width=27).grid(row=3, column=1, padx=10, pady=5)

        def save():
            with connect() as db:
                cls = db.execute("SELECT id FROM classes WHERE name=?", (class_var.get(),)).fetchone()
                try:
                    db.execute("UPDATE subjects SET name=?,teacher_name=?,class_id=?,active=? WHERE id=?", (name_var.get().strip(), teacher_var.get().strip(), cls["id"] if cls else None, 1 if active_var.get() == "Active" else 0, subject_id))
                except sqlite3.IntegrityError:
                    messagebox.showerror("Modification impossible", "Cette matière existe déjà dans cette classe.", parent=dialog)
                    return
                db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('modification','matière',?,?)", (subject_id, f"{name_var.get().strip()}; {class_var.get()}"))
            dialog.destroy()
            self.refresh_academics()

        ttk.Button(dialog, text="Enregistrer", command=save, style="Accent.TButton").grid(row=4, column=0, columnspan=2, pady=8)

    def delete_selected_subject(self) -> None:
        selected = self.subject_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez une matière dans la liste.")
            return
        subject_id = int(selected[0])
        with connect() as db:
            row = db.execute("SELECT name,active FROM subjects WHERE id=?", (subject_id,)).fetchone()
        if not row:
            return
        action = "archiver" if row["active"] else "réactiver"
        if not messagebox.askyesno("Confirmer", f"Voulez-vous {action} la matière « {row['name']} » ? Les notes existantes seront conservées.", icon="warning"):
            return
        active = 0 if row["active"] else 1
        with connect() as db:
            db.execute("UPDATE subjects SET active=? WHERE id=?", (active, subject_id))
            db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES(?,?,?,?)", ("archivage" if not active else "réactivation", "matière", subject_id, row["name"]))
            return
        self.refresh_academics()

    def refresh_grade_subjects(self, _event=None) -> None:
        selection = self.grade_student.get()
        try:
            student_id = int(selection.split(" — ", 1)[0])
        except (ValueError, IndexError):
            self.grade_subject["values"] = []
            self.grade_subject.set("")
            self.refresh_grade_list()
            return
        with connect() as db:
            row = db.execute("SELECT class_id FROM students WHERE id=?", (student_id,)).fetchone()
            subjects = [] if not row or row["class_id"] is None else db.execute(
                "SELECT id,name,teacher_name FROM subjects WHERE class_id=? AND active=1 ORDER BY name", (row["class_id"],)
            ).fetchall()
        choices = [f"{item['id']} — {item['name']}" + (f" — {item['teacher_name']}" if item["teacher_name"] else "") for item in subjects]
        self.grade_subject["values"] = choices
        if self.grade_subject.get() not in choices:
            self.grade_subject.set(choices[0] if choices else "")
        self.refresh_grade_list()

    def refresh_grade_list(self) -> None:
        self.clear(self.grade_tree)
        selection = self.grade_student.get()
        period = self.grade_period.get()
        year = self.grade_year.get().strip()
        try:
            student_id = int(selection.split(" — ", 1)[0])
        except (ValueError, IndexError):
            self.grade_average.configure(text="Moyenne générale : -")
            return
        with connect() as db:
            rows = db.execute("""SELECT g.id,s.name,s.teacher_name,g.score,g.max_score,g.comment
                FROM grades g JOIN subjects s ON s.id=g.subject_id
                WHERE g.student_id=? AND g.period=? AND g.school_year=? ORDER BY s.name""", (student_id, period, year)).fetchall()
        normalized = []
        for row in rows:
            score10 = row["score"] * 10 / row["max_score"]
            normalized.append(score10)
            self.grade_tree.insert("", "end", iid=str(row["id"]), values=(row["name"], row["teacher_name"], row["score"], row["max_score"], f"{score10:.2f} / 10", row["comment"]))
        if normalized:
            average = sum(normalized) / len(normalized)
            self.grade_average.configure(text=f"Moyenne générale : {average:.2f} / 10 ({len(normalized)} matière(s))")
        else:
            self.grade_average.configure(text="Moyenne générale : aucune note pour cette période")

    def add_grade(self) -> None:
        try:
            student_id = int(self.grade_student.get().split(" — ", 1)[0])
            subject_id = int(self.grade_subject.get().split(" — ", 1)[0])
            score = float(self.grade_score.get().strip().replace(",", "."))
            max_score = float(self.grade_max.get().strip().replace(",", "."))
            period = self.grade_period.get()
            year = self.grade_year.get().strip()
            if max_score <= 0 or score < 0 or score > max_score or not period or not year:
                raise ValueError
        except (ValueError, IndexError):
            messagebox.showwarning("Note invalide", "Choisissez l'élève, la matière, la période et indiquez une note comprise entre 0 et la note maximale.")
            return
        comment = self.grade_comment.get().strip()
        with connect() as db:
            existing = db.execute("""SELECT id,score,max_score FROM grades
                WHERE student_id=? AND subject_id=? AND school_year=? AND period=?""", (student_id, subject_id, year, period)).fetchone()
            if existing:
                if not messagebox.askyesno("Corriger la note", f"Une note existe déjà ({existing['score']} / {existing['max_score']}). Voulez-vous la remplacer ?"):
                    return
                db.execute("UPDATE grades SET score=?,max_score=?,comment=? WHERE id=?", (score, max_score, comment, existing["id"]))
                action, grade_id = "correction", existing["id"]
                details = f"{period} {year}; {existing['score']} -> {score}"
            else:
                cur = db.execute("""INSERT INTO grades(student_id,subject_id,school_year,period,score,max_score,comment)
                    VALUES(?,?,?,?,?,?,?)""", (student_id, subject_id, year, period, score, max_score, comment))
                action, grade_id = "création", cur.lastrowid
                details = f"{period} {year}; note {score}/{max_score}"
            db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES(?,?,?,?)", (action, "note", grade_id, details))
        self.grade_score.set("")
        self.grade_comment.set("")
        self.refresh_grade_list()

    def export_bulletin(self) -> None:
        try:
            student_id = int(self.grade_student.get().split(" — ", 1)[0])
        except (ValueError, IndexError):
            messagebox.showinfo("Choisir un élève", "Sélectionnez un élève pour préparer son bulletin.")
            return
        period, year = self.grade_period.get(), self.grade_year.get().strip()
        with connect() as db:
            student = db.execute("""SELECT s.matricule,s.surname,s.given_name,c.name AS class_name
                FROM students s LEFT JOIN classes c ON c.id=s.class_id WHERE s.id=?""", (student_id,)).fetchone()
            grades = db.execute("""SELECT sub.name,sub.teacher_name,g.score,g.max_score,g.comment
                FROM grades g JOIN subjects sub ON sub.id=g.subject_id
                WHERE g.student_id=? AND g.period=? AND g.school_year=? ORDER BY sub.name""", (student_id, period, year)).fetchall()
        if not grades:
            messagebox.showinfo("Aucune note", "Aucune note n'est enregistrée pour cet élève et cette période.")
            return
        try:
            from reportlab.lib import colors
            from reportlab.lib.enums import TA_CENTER
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import mm
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        except ImportError:
            messagebox.showerror("Composant absent", "Installez les dépendances avec : python -m pip install -r requirements.txt")
            return
        from html import escape

        bulletin_dir = DATA_DIR / "bulletins"
        bulletin_dir.mkdir(parents=True, exist_ok=True)
        default_name = f"bulletin-{student['matricule']}-{year}-{period.lower().replace(' ', '-')}.pdf"
        target = filedialog.asksaveasfilename(
            title="Enregistrer le bulletin PDF", initialdir=bulletin_dir, initialfile=default_name,
            defaultextension=".pdf", filetypes=[("Document PDF", "*.pdf")],
        )
        if not target:
            return
        body_font, bold_font = "Helvetica", "Helvetica-Bold"
        font_path, bold_path = Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")
        if font_path.exists():
            pdfmetrics.registerFont(TTFont("SchoolArial", str(font_path)))
            body_font = "SchoolArial"
        if bold_path.exists():
            pdfmetrics.registerFont(TTFont("SchoolArialBold", str(bold_path)))
            bold_font = "SchoolArialBold"
        try:
            import arabic_reshaper
            from bidi.algorithm import get_display

            def prepare_text(value):
                value = str(value or "")
                if any("\u0600" <= char <= "\u06ff" for char in value):
                    value = get_display(arabic_reshaper.reshape(value))
                return escape(value)
        except ImportError:
            prepare_text = lambda value: escape(str(value or ""))

        try:
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

            styles = getSampleStyleSheet()
            body_style = ParagraphStyle("BulletinBody", parent=styles["Normal"], fontName=body_font, fontSize=9, leading=12)
            centered_style = ParagraphStyle("BulletinTitle", parent=body_style, fontName=bold_font, fontSize=15, leading=20, alignment=TA_CENTER, textColor=colors.HexColor("#16324F"))
            school_style = ParagraphStyle("SchoolName", parent=body_style, fontName=bold_font, fontSize=12, leading=16, alignment=TA_CENTER)
            story = [
                Paragraph(prepare_text("Medersa Franco-Arabe Sabiloul Houda"), school_style),
                Paragraph(prepare_text("Mamaribougou - Mali"), body_style),
                Spacer(1, 5 * mm),
                Paragraph(prepare_text("BULLETIN DE NOTES"), centered_style),
                Paragraph(prepare_text(f"{period} - Année scolaire {year}"), body_style),
                Spacer(1, 4 * mm),
            ]
            student_table = Table([
                ["Élève", prepare_text(f"{student['surname']} {student['given_name']}"), "Matricule", prepare_text(student["matricule"])],
                ["Classe", prepare_text(student["class_name"] or "Non renseignée"), "Période", prepare_text(period)],
            ], colWidths=[23 * mm, 70 * mm, 25 * mm, 62 * mm])
            student_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF0F4")),
                ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#EAF0F4")),
                ("FONTNAME", (0, 0), (0, -1), bold_font),
                ("FONTNAME", (2, 0), (2, -1), bold_font),
                ("FONTNAME", (1, 0), (1, -1), body_font),
                ("FONTNAME", (3, 0), (3, -1), body_font),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCD6DE")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]))
            story.extend([student_table, Spacer(1, 7 * mm)])
            table_data = [["Matière", "Enseignant", "Note", "Sur", "Équiv. /10", "Appréciation"]]
            normalized_scores = []
            for grade in grades:
                normalized_score = grade["score"] * 10 / grade["max_score"]
                normalized_scores.append(normalized_score)
                table_data.append([
                    Paragraph(prepare_text(grade["name"]), body_style),
                    Paragraph(prepare_text(grade["teacher_name"]), body_style),
                    f"{grade['score']:g}", f"{grade['max_score']:g}",
                    f"{normalized_score:.2f}", Paragraph(prepare_text(grade["comment"]), body_style),
                ])
            marks_table = Table(table_data, colWidths=[42 * mm, 36 * mm, 20 * mm, 18 * mm, 27 * mm, 37 * mm], repeatRows=1)
            marks_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), bold_font),
                ("FONTNAME", (0, 1), (-1, -1), body_font),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCD6DE")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (2, 1), (4, -1), "CENTER"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8FA")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]))
            average = sum(normalized_scores) / len(normalized_scores)
            if average >= 8:
                mention = "Très bien"
            elif average >= 7:
                mention = "Bien"
            elif average >= 6:
                mention = "Assez bien"
            elif average >= 5:
                mention = "Passable"
            else:
                mention = "Insuffisant"
            average_style = ParagraphStyle("Average", parent=body_style, fontName=bold_font, fontSize=12, textColor=colors.HexColor("#168C85"))
            story.extend([
                marks_table,
                Spacer(1, 6 * mm),
                Paragraph(prepare_text(f"Moyenne générale : {average:.2f} / 10 - Mention : {mention}"), average_style),
                Spacer(1, 18 * mm),
                Table([["Signature du directeur", "Signature du parent"]], colWidths=[90 * mm, 90 * mm], style=TableStyle([
                    ("LINEABOVE", (0, 0), (0, 0), 0.6, colors.HexColor("#81909B")),
                    ("LINEABOVE", (1, 0), (1, 0), 0.6, colors.HexColor("#81909B")),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("FONTNAME", (0, 0), (-1, -1), body_font),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                ])),
            ])
            output = Path(target)
            document = SimpleDocTemplate(
                str(output), pagesize=A4, rightMargin=15 * mm, leftMargin=15 * mm,
                topMargin=14 * mm, bottomMargin=16 * mm, title=f"Bulletin {student['matricule']} {year} {period}",
            )
            document.build(story)
        except Exception as exc:
            messagebox.showerror("Création du bulletin impossible", str(exc))
            return
        try:
            if sys.platform == "win32":
                os.startfile(output)
            elif sys.platform == "darwin":
                subprocess.Popen(["open", str(output)])
            else:
                subprocess.Popen(["xdg-open", str(output)])
        except OSError:
            messagebox.showinfo("Bulletin créé", f"Le PDF a été enregistré ici :\n{output}")

    def refresh_attendance_table(self) -> None:
        class_name = self.attendance_class.get()
        date_text = self.attendance_date.get().strip()
        try:
            datetime.strptime(date_text, "%Y-%m-%d")
        except ValueError:
            self.clear(self.attendance_tree)
            return
        with connect() as db:
            rows = db.execute("""SELECT a.id,a.attendance_date,s.matricule,s.surname,s.given_name,a.status,a.note
                FROM attendance a JOIN students s ON s.id=a.student_id
                JOIN classes c ON c.id=s.class_id
                WHERE a.attendance_date=? AND c.name=? ORDER BY s.surname,s.given_name""", (date_text, class_name)).fetchall()
        self.clear(self.attendance_tree)
        for row in rows:
            self.attendance_tree.insert("", "end", iid=str(row["id"]), values=(row["attendance_date"], row["matricule"], f"{row['surname']} {row['given_name']}", row["status"], row["note"]))

    def add_attendance(self) -> None:
        selection = self.attendance_student.get()
        status = self.attendance_status.get()
        date_text = self.attendance_date.get().strip()
        try:
            student_id = int(selection.split(" — ", 1)[0])
            datetime.strptime(date_text, "%Y-%m-%d")
            if status not in ("Présent", "Absent", "Retard"):
                raise ValueError
        except (ValueError, IndexError):
            messagebox.showwarning("Présence invalide", "Vérifiez la date, la classe, l'élève et le statut.")
            return
        note = self.attendance_note.get().strip()
        with connect() as db:
            old = db.execute("SELECT id,status,note FROM attendance WHERE student_id=? AND attendance_date=?", (student_id, date_text)).fetchone()
            if old:
                if old["status"] == status and old["note"] == note:
                    messagebox.showinfo("Déjà enregistré", "Cette présence a déjà été enregistrée avec les mêmes informations.")
                    return
                if not messagebox.askyesno("Corriger la présence", "Une présence existe déjà pour cet élève à cette date. Voulez-vous la corriger ?"):
                    return
                db.execute("UPDATE attendance SET status=?,note=? WHERE id=?", (status, note, old["id"]))
                action = "correction"
                attendance_id = old["id"]
                detail = f"{date_text}; {old['status']} -> {status}"
            else:
                cur = db.execute("INSERT INTO attendance(student_id,attendance_date,status,note) VALUES(?,?,?,?)", (student_id, date_text, status, note))
                action = "création"
                attendance_id = cur.lastrowid
                detail = f"{date_text}; {status}"
            db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES(?,?,?,?)", (action, "présence", attendance_id, detail))
        self.attendance_note.set("")
        self.refresh_attendance_table()

    def edit_selected_attendance(self) -> None:
        selected = self.attendance_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez une présence dans la liste.")
            return
        attendance_id = int(selected[0])
        with connect() as db:
            row = db.execute("SELECT attendance_date,status,note,student_id FROM attendance WHERE id=?", (attendance_id,)).fetchone()
        if not row:
            return

        def save(values, dialog):
            try:
                datetime.strptime(values["date"], "%Y-%m-%d")
                if values["status"] not in ("Présent", "Absent", "Retard"):
                    raise ValueError
            except ValueError:
                messagebox.showwarning("Présence invalide", "Indiquez une date au format AAAA-MM-JJ et un statut valide.", parent=dialog)
                return False
            try:
                with connect() as db:
                    db.execute("UPDATE attendance SET attendance_date=?,status=?,note=? WHERE id=?", (values["date"], values["status"], values["note"], attendance_id))
                    db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('modification','présence',?,?)", (attendance_id, f"{values['date']}; {values['status']}"))
            except sqlite3.IntegrityError:
                messagebox.showerror("Présence déjà enregistrée", "Cet élève a déjà une présence enregistrée à cette date.", parent=dialog)
                return False
            self.attendance_date.set(values["date"])
            self.refresh_attendance_table()
            return True

        fields = [
            {"label": "Date (AAAA-MM-JJ)", "key": "date", "value": row["attendance_date"]},
            {"label": "Statut", "key": "status", "value": row["status"], "choices": ["Présent", "Absent", "Retard"]},
            {"label": "Remarque", "key": "note", "value": row["note"]},
        ]
        self._edit_dialog("Modifier la présence", fields, save)

    def delete_selected_attendance(self) -> None:
        selected = self.attendance_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez une présence dans la liste.")
            return
        attendance_id = int(selected[0])
        if not messagebox.askyesno("Confirmer la suppression", "Supprimer cette présence ?", icon="warning"):
            return
        with connect() as db:
            row = db.execute("SELECT attendance_date,status FROM attendance WHERE id=?", (attendance_id,)).fetchone()
            db.execute("DELETE FROM attendance WHERE id=?", (attendance_id,))
            db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('suppression','présence',?,?)", (attendance_id, f"{row['attendance_date']}; {row['status']}" if row else ""))
        self.refresh_attendance_table()

    def add_fee(self) -> None:
        selection = self.fee_student.get()
        try:
            student_id = int(selection.split(" — ", 1)[0])
            amount = int(self.fee_amount.get().replace(" ", ""))
            label = self.fee_label.get().strip()
            year = self.fee_year.get().strip()
            if amount <= 0 or not label or not year:
                raise ValueError
        except (ValueError, IndexError):
            messagebox.showwarning("Frais incomplet", "Sélectionnez un élève, indiquez l'année, le nom du frais et un montant positif en FCFA.")
            return
        try:
            with connect() as db:
                db.execute("INSERT INTO fees(student_id,school_year,label,amount) VALUES(?,?,?,?)", (student_id, year, label, amount))
                fee_id = db.execute("SELECT last_insert_rowid() AS id").fetchone()["id"]
                db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('création','frais',?,?)", (fee_id, f"{label}; {amount} FCFA; {year}"))
        except sqlite3.IntegrityError:
            messagebox.showwarning("Frais déjà existant", "Un frais avec ce nom existe déjà pour cet élève et cette année scolaire.")
            return
        except sqlite3.Error as exc:
            messagebox.showerror("Erreur", str(exc))
            return
        self.fee_amount.set("")
        self.fee_label.set("")
        self.refresh_all()

    def edit_selected_fee(self) -> None:
        selected = self.fee_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez un frais dans la liste.")
            return
        fee_id = int(selected[0])
        with connect() as db:
            row = db.execute("SELECT school_year,label,amount,COALESCE((SELECT SUM(amount) FROM payments WHERE fee_id=fees.id),0) AS paid FROM fees WHERE id=?", (fee_id,)).fetchone()
        if not row:
            return

        def save(values, dialog):
            try:
                amount = int(values["amount"].replace(" ", ""))
                if not values["year"] or not values["label"] or amount <= 0:
                    raise ValueError
                if amount < row["paid"]:
                    messagebox.showwarning("Montant trop faible", f"Le montant dû ne peut pas être inférieur aux versements déjà reçus ({row['paid']:,} FCFA).".replace(",", " "), parent=dialog)
                    return False
            except ValueError:
                messagebox.showwarning("Frais invalide", "L'année, le libellé et un montant positif sont requis.", parent=dialog)
                return False
            try:
                with connect() as db:
                    db.execute("UPDATE fees SET school_year=?,label=?,amount=? WHERE id=?", (values["year"], values["label"], amount, fee_id))
                    db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('modification','frais',?,?)", (fee_id, f"{values['label']}; {amount} FCFA; {values['year']}"))
            except sqlite3.IntegrityError:
                messagebox.showerror("Frais déjà existant", "Ce frais existe déjà pour cet élève et cette année.", parent=dialog)
                return False
            self.refresh_all()
            return True

        self._edit_dialog("Modifier les frais", [
            {"label": "Année scolaire", "key": "year", "value": row["school_year"]},
            {"label": "Frais / mensualité", "key": "label", "value": row["label"]},
            {"label": "Montant dû (FCFA)", "key": "amount", "value": row["amount"]},
        ], save)

    def delete_selected_fee(self) -> None:
        selected = self.fee_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez un frais dans la liste.")
            return
        fee_id = int(selected[0])
        if not messagebox.askyesno("Confirmer la suppression", "Supprimer ce frais ? Les paiements déjà enregistrés seront conservés dans le journal des recettes, mais détachés de ce frais.", icon="warning"):
            return
        with connect() as db:
            row = db.execute("SELECT label FROM fees WHERE id=?", (fee_id,)).fetchone()
            db.execute("UPDATE payments SET fee_id=NULL WHERE fee_id=?", (fee_id,))
            db.execute("DELETE FROM fees WHERE id=?", (fee_id,))
            db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('suppression','frais',?,?)", (fee_id, row["label"] if row else ""))
        self.refresh_all()

    def update_payment_balance(self, _event=None) -> None:
        selection = self.payment_fee.get()
        try:
            fee_id = int(selection.split(" — ", 1)[0])
        except (ValueError, IndexError):
            self.payment_balance.configure(text="Reste à payer : -")
            return
        with connect() as db:
            row = db.execute("""SELECT f.amount-COALESCE(SUM(p.amount),0) AS remaining
                FROM fees f LEFT JOIN payments p ON p.fee_id=f.id WHERE f.id=? GROUP BY f.id""", (fee_id,)).fetchone()
        remaining = row["remaining"] if row else 0
        self.payment_balance.configure(text=f"Reste à payer : {remaining:,} FCFA".replace(",", " "))

    def add_payment(self) -> None:
        selection = self.payment_fee.get()
        try:
            fee_id = int(selection.split(" — ", 1)[0])
            amount = int(self.payment_amount.get().replace(" ", ""))
            if amount <= 0:
                raise ValueError
        except (ValueError, IndexError):
            messagebox.showwarning("Paiement invalide", "Choisissez un frais dû et saisissez un montant positif en FCFA.")
            return
        from datetime import date
        with connect() as db:
            fee = db.execute("""SELECT f.student_id, f.label, f.amount-COALESCE(SUM(p.amount),0) AS remaining
                FROM fees f LEFT JOIN payments p ON p.fee_id=f.id WHERE f.id=? GROUP BY f.id""", (fee_id,)).fetchone()
            if fee is None:
                messagebox.showerror("Frais introuvable", "Actualisez l'écran et choisissez un frais existant.")
                return
            if amount > fee["remaining"]:
                messagebox.showwarning("Montant supérieur au reste dû", f"Il reste {fee['remaining']:,} FCFA à payer pour ce frais.".replace(",", " "))
                return
            next_id = db.execute("SELECT COALESCE(MAX(id),0)+1 AS next_id FROM payments").fetchone()["next_id"]
            receipt_no = f"SH-{date.today():%y%m%d}-{next_id:04d}"
            cur = db.execute("""INSERT INTO payments(receipt_no,student_id,fee_id,paid_at,amount,reason)
                VALUES(?,?,?,?,?,?)""", (receipt_no, fee["student_id"], fee_id, date.today().isoformat(), amount, fee["label"]))
            db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('création','paiement',?,?)", (cur.lastrowid, f"{receipt_no}; {amount} FCFA"))
        self.payment_amount.set("")
        self.refresh_all()

    def edit_selected_payment(self) -> None:
        selected = self.payment_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez un paiement dans la liste.")
            return
        payment_id = int(selected[0])
        with connect() as db:
            row = db.execute("SELECT paid_at,amount,reason,note,receipt_no FROM payments WHERE id=?", (payment_id,)).fetchone()
        if not row:
            return

        def save(values, dialog):
            try:
                datetime.strptime(values["date"], "%Y-%m-%d")
                amount = int(values["amount"].replace(" ", ""))
                if amount <= 0 or not values["reason"]:
                    raise ValueError
            except ValueError:
                messagebox.showwarning("Paiement invalide", "La date doit être au format AAAA-MM-JJ, le motif requis et le montant positif.", parent=dialog)
                return False
            with connect() as db:
                db.execute("UPDATE payments SET paid_at=?,amount=?,reason=?,note=? WHERE id=?", (values["date"], amount, values["reason"], values["note"], payment_id))
                db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('modification','paiement',?,?)", (payment_id, f"{row['receipt_no']}; {amount} FCFA"))
            self.refresh_all()
            return True

        self._edit_dialog("Modifier le paiement", [
            {"label": "Date (AAAA-MM-JJ)", "key": "date", "value": row["paid_at"]},
            {"label": "Montant versé (FCFA)", "key": "amount", "value": row["amount"]},
            {"label": "Motif", "key": "reason", "value": row["reason"]},
            {"label": "Remarque", "key": "note", "value": row["note"]},
        ], save)

    def delete_selected_payment(self) -> None:
        selected = self.payment_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez un paiement dans la liste.")
            return
        payment_id = int(selected[0])
        if not messagebox.askyesno("Confirmer la suppression", "Supprimer ce paiement ? Le solde des frais sera recalculé et le reçu correspondant ne sera plus valide.", icon="warning"):
            return
        with connect() as db:
            row = db.execute("SELECT receipt_no,amount FROM payments WHERE id=?", (payment_id,)).fetchone()
            db.execute("DELETE FROM payments WHERE id=?", (payment_id,))
            db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('suppression','paiement',?,?)", (payment_id, f"{row['receipt_no']}; {row['amount']} FCFA" if row else ""))
        self.refresh_all()

    def export_selected_receipt(self) -> None:
        selected = self.payment_tree.selection()
        if not selected:
            messagebox.showinfo("Choisir un paiement", "Sélectionnez d'abord un paiement dans la liste.")
            return
        receipt_no = self.payment_tree.item(selected[0], "values")[0]
        with connect() as db:
            receipt = db.execute("""SELECT p.receipt_no,p.paid_at,p.amount,p.reason,
                s.matricule,s.surname,s.given_name,s.parent_name,c.name AS class_name,
                f.label AS fee_label,f.school_year,f.amount AS fee_amount,
                (SELECT COALESCE(SUM(p2.amount),0) FROM payments p2
                 WHERE p2.fee_id=f.id AND p2.id<=p.id) AS paid_to_date
                FROM payments p JOIN students s ON s.id=p.student_id
                LEFT JOIN classes c ON c.id=s.class_id
                LEFT JOIN fees f ON f.id=p.fee_id
                WHERE p.receipt_no=?""", (receipt_no,)).fetchone()
        if receipt is None:
            messagebox.showerror("Reçu introuvable", "Ce paiement n'existe plus dans la base de données.")
            return
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.units import mm
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.pdfgen import canvas
        except ImportError:
            messagebox.showerror("Composant absent", "Installez les dépendances PDF avec : pip install -r requirements.txt")
            return

        default_dir = DATA_DIR / "recus"
        default_dir.mkdir(parents=True, exist_ok=True)
        target = filedialog.asksaveasfilename(
            title="Enregistrer le reçu PDF",
            initialdir=default_dir,
            initialfile=f"recu-{receipt_no}.pdf",
            defaultextension=".pdf",
            filetypes=[("Document PDF", "*.pdf")],
        )
        if not target:
            return
        output = Path(target)

        body_font, bold_font = "Helvetica", "Helvetica-Bold"
        font_path = Path("C:/Windows/Fonts/arial.ttf")
        bold_path = Path("C:/Windows/Fonts/arialbd.ttf")
        if font_path.exists():
            pdfmetrics.registerFont(TTFont("SchoolArial", str(font_path)))
            body_font = "SchoolArial"
        if bold_path.exists():
            pdfmetrics.registerFont(TTFont("SchoolArialBold", str(bold_path)))
            bold_font = "SchoolArialBold"
        try:
            import arabic_reshaper
            from bidi.algorithm import get_display

            def pdf_text(value):
                value = str(value or "")
                if any("\u0600" <= char <= "\u06ff" for char in value):
                    return get_display(arabic_reshaper.reshape(value))
                return value
        except ImportError:
            pdf_text = lambda value: str(value or "")

        try:
            pdf = canvas.Canvas(str(output), pagesize=A4)
            width, height = A4
            navy = colors.HexColor("#16324F")
            teal = colors.HexColor("#168C85")
            pale = colors.HexColor("#F0F5F8")
            pdf.setFillColor(navy)
            pdf.roundRect(18 * mm, height - 54 * mm, width - 36 * mm, 34 * mm, 4 * mm, fill=1, stroke=0)
            pdf.setFillColor(colors.white)
            pdf.setFont(bold_font, 17)
            pdf.drawString(26 * mm, height - 33 * mm, pdf_text("Medersa Franco-Arabe Sabiloul Houda"))
            pdf.setFont(body_font, 10)
            pdf.drawString(26 * mm, height - 42 * mm, pdf_text("Mamaribougou - Mali"))
            pdf.setFillColor(teal)
            pdf.roundRect(width - 67 * mm, height - 69 * mm, 49 * mm, 13 * mm, 3 * mm, fill=1, stroke=0)
            pdf.setFillColor(colors.white)
            pdf.setFont(bold_font, 11)
            pdf.drawCentredString(width - 42.5 * mm, height - 64 * mm, pdf_text("REÇU DE PAIEMENT"))

            y = height - 86 * mm
            pdf.setFillColor(navy)
            pdf.setFont(bold_font, 11)
            pdf.drawString(22 * mm, y, pdf_text("Numéro du reçu"))
            pdf.drawRightString(width - 22 * mm, y, pdf_text("Date : " + receipt["paid_at"]))
            pdf.setFont(body_font, 12)
            pdf.drawString(22 * mm, y - 8 * mm, pdf_text(receipt["receipt_no"]))

            y -= 24 * mm
            pdf.setFillColor(pale)
            pdf.roundRect(18 * mm, y - 36 * mm, width - 36 * mm, 43 * mm, 3 * mm, fill=1, stroke=0)
            pdf.setFillColor(navy)
            pdf.setFont(bold_font, 11)
            pdf.drawString(24 * mm, y, pdf_text("Élève"))
            pdf.setFont(body_font, 11)
            pdf.drawString(24 * mm, y - 7 * mm, pdf_text(f"{receipt['surname']} {receipt['given_name']}"))
            pdf.setFont(bold_font, 10)
            pdf.drawString(24 * mm, y - 17 * mm, pdf_text("Matricule"))
            pdf.setFont(body_font, 10)
            pdf.drawString(24 * mm, y - 24 * mm, pdf_text(receipt["matricule"]))
            pdf.setFont(bold_font, 10)
            pdf.drawString(105 * mm, y, pdf_text("Classe"))
            pdf.setFont(body_font, 11)
            pdf.drawString(105 * mm, y - 7 * mm, pdf_text(receipt["class_name"] or "Non renseignée"))
            pdf.setFont(bold_font, 10)
            pdf.drawString(105 * mm, y - 17 * mm, pdf_text("Responsable"))
            pdf.setFont(body_font, 10)
            pdf.drawString(105 * mm, y - 24 * mm, pdf_text(receipt["parent_name"] or "Non renseigné"))

            y -= 53 * mm
            pdf.setFont(bold_font, 11)
            pdf.setFillColor(navy)
            pdf.drawString(22 * mm, y, pdf_text("Détail du règlement"))
            pdf.setStrokeColor(colors.HexColor("#D8E1E8"))
            pdf.line(22 * mm, y - 3 * mm, width - 22 * mm, y - 3 * mm)
            pdf.setFont(body_font, 11)
            label = receipt["fee_label"] or receipt["reason"]
            if receipt["school_year"]:
                label += f" ({receipt['school_year']})"
            pdf.drawString(24 * mm, y - 13 * mm, pdf_text("Motif : " + label))
            pdf.drawString(24 * mm, y - 22 * mm, pdf_text("Montant reçu :"))
            pdf.setFont(bold_font, 16)
            pdf.setFillColor(teal)
            pdf.drawRightString(width - 24 * mm, y - 22 * mm, f"{receipt['amount']:,} FCFA".replace(",", " "))
            if receipt["fee_amount"] is not None:
                paid_to_date = receipt["paid_to_date"] or 0
                remaining = max(0, receipt["fee_amount"] - paid_to_date)
                pdf.setFillColor(navy)
                pdf.setFont(body_font, 10)
                pdf.drawString(24 * mm, y - 34 * mm, pdf_text(f"Total du frais : {receipt['fee_amount']:,} FCFA".replace(",", " ")))
                pdf.drawRightString(width - 24 * mm, y - 34 * mm, pdf_text(f"Reste dû après ce versement : {remaining:,} FCFA".replace(",", " ")))

            pdf.setStrokeColor(colors.HexColor("#AAB8C4"))
            pdf.line(22 * mm, 45 * mm, 83 * mm, 45 * mm)
            pdf.line(width - 83 * mm, 45 * mm, width - 22 * mm, 45 * mm)
            pdf.setFillColor(navy)
            pdf.setFont(body_font, 9)
            pdf.drawCentredString(52 * mm, 39 * mm, pdf_text("Signature du caissier"))
            pdf.drawCentredString(width - 52 * mm, 39 * mm, pdf_text("Signature du parent"))
            pdf.setFillColor(colors.HexColor("#687888"))
            pdf.setFont(body_font, 8)
            pdf.drawCentredString(width / 2, 20 * mm, pdf_text("Document généré par le logiciel scolaire Sabiloul Houda"))
            pdf.save()
        except Exception as exc:
            messagebox.showerror("Création du reçu impossible", str(exc))
            return

        try:
            if sys.platform == "win32":
                os.startfile(output)
            elif sys.platform == "darwin":
                subprocess.Popen(["open", str(output)])
            else:
                subprocess.Popen(["xdg-open", str(output)])
        except OSError:
            messagebox.showinfo("Reçu créé", f"Le PDF a été enregistré ici :\n{output}")

    def add_expense(self) -> None:
        data = {k: v.get().strip() for k, v in self.expense_vars.items()}
        try:
            amount = int(data["amount"].replace(" ", ""))
            if amount <= 0 or not all(data[k] for k in ("date", "category", "description")):
                raise ValueError
        except ValueError:
            messagebox.showwarning("Dépense invalide", "Date, catégorie, description et montant positif sont requis.")
            return
        with connect() as db:
            cur = db.execute("INSERT INTO expenses(spent_at,category,description,amount,payee) VALUES(?,?,?,?,?)", (data["date"], data["category"], data["description"], amount, data["payee"]))
            db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('création','dépense',?,?)", (cur.lastrowid, f"{amount} FCFA"))
        for var in self.expense_vars.values():
            var.set("")
        self.refresh_all()

    def edit_selected_expense(self) -> None:
        selected = self.expense_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez d'abord une dépense.")
            return
        expense_id = int(selected[0])
        with connect() as db:
            row = db.execute("SELECT spent_at,category,description,amount,payee FROM expenses WHERE id=?", (expense_id,)).fetchone()
        if not row:
            return
        dialog = tk.Toplevel(self)
        dialog.title("Modifier la dépense")
        dialog.transient(self)
        dialog.grab_set()
        labels = [("Date", "spent_at"), ("Catégorie", "category"), ("Description", "description"), ("Montant FCFA", "amount"), ("Bénéficiaire", "payee")]
        variables = {}
        for idx, (label, key) in enumerate(labels):
            ttk.Label(dialog, text=label).grid(row=idx, column=0, padx=10, pady=4, sticky="w")
            variables[key] = tk.StringVar(value=str(row[key] or ""))
            ttk.Entry(dialog, textvariable=variables[key], width=34).grid(row=idx, column=1, padx=10, pady=4)

        def save():
            values = {key: var.get().strip() for key, var in variables.items()}
            try:
                amount = int(values["amount"].replace(" ", ""))
                if amount <= 0 or not values["spent_at"] or not values["category"] or not values["description"]:
                    raise ValueError
            except ValueError:
                messagebox.showwarning("Dépense invalide", "Date, catégorie, description et montant positif sont requis.", parent=dialog)
                return
            with connect() as db:
                db.execute("UPDATE expenses SET spent_at=?,category=?,description=?,amount=?,payee=? WHERE id=?", (values["spent_at"], values["category"], values["description"], amount, values["payee"], expense_id))
                db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('modification','dépense',?,?)", (expense_id, f"{amount} FCFA"))
            dialog.destroy()
            self.refresh_all()

        ttk.Button(dialog, text="Enregistrer", command=save, style="Accent.TButton").grid(row=len(labels), column=0, columnspan=2, pady=8)

    def delete_selected_expense(self) -> None:
        selected = self.expense_tree.selection()
        if not selected:
            messagebox.showinfo("Sélection nécessaire", "Sélectionnez d'abord une dépense.")
            return
        expense_id = int(selected[0])
        if not messagebox.askyesno("Confirmer la suppression", "Supprimer cette dépense du journal comptable ?", icon="warning"):
            return
        with connect() as db:
            row = db.execute("SELECT description FROM expenses WHERE id=?", (expense_id,)).fetchone()
            db.execute("DELETE FROM expenses WHERE id=?", (expense_id,))
            db.execute("INSERT INTO audit_log(action,entity,entity_id,details) VALUES('suppression','dépense',?,?)", (expense_id, row["description"] if row else ""))
        self.refresh_all()

    def load_report(self) -> None:
        report = self.report_choice.get()
        queries = {
            "Liste des élèves": (
                ["Matricule", "Nom", "Prénom", "Sexe", "Naissance", "Classe", "Inscription", "Responsable", "Téléphone", "Statut"],
                """SELECT s.matricule,s.surname,s.given_name,s.sex,s.birth_date,
                    COALESCE(c.name,''),s.enrollment_date,s.parent_name,s.parent_phone,s.status
                    FROM students s LEFT JOIN classes c ON c.id=s.class_id
                    WHERE (:class_name IS NULL OR c.name=:class_name)
                    ORDER BY s.surname,s.given_name""",
            ),
            "Élèves par classe (A à Z)": (
                ["Classe", "Matricule", "Nom", "Prénom", "Sexe", "Téléphone", "Statut"],
                """SELECT COALESCE(c.name,''),s.matricule,s.surname,s.given_name,s.sex,s.parent_phone,s.status
                    FROM students s LEFT JOIN classes c ON c.id=s.class_id
                    WHERE (:class_name IS NULL OR c.name=:class_name)
                    ORDER BY c.name COLLATE NOCASE,s.surname COLLATE NOCASE,s.given_name COLLATE NOCASE""",
            ),
            "Matières par classe": (
                ["Classe", "Matière", "Enseignant"],
                """SELECT c.name,sub.name,sub.teacher_name FROM subjects sub
                    JOIN classes c ON c.id=sub.class_id
                    WHERE sub.active=1 AND (:class_name IS NULL OR c.name=:class_name)
                    ORDER BY c.name COLLATE NOCASE,sub.name COLLATE NOCASE""",
            ),
            "Frais impayés": (
                ["Matricule", "Nom", "Prénom", "Classe", "Année scolaire", "Frais / mois", "Dû (FCFA)", "Payé (FCFA)", "Reste (FCFA)"],
                """SELECT s.matricule,s.surname,s.given_name,COALESCE(c.name,''),f.school_year,f.label,
                    f.amount,COALESCE(SUM(p.amount),0),f.amount-COALESCE(SUM(p.amount),0)
                    FROM fees f JOIN students s ON s.id=f.student_id
                    LEFT JOIN classes c ON c.id=s.class_id
                    LEFT JOIN payments p ON p.fee_id=f.id
                    WHERE (:class_name IS NULL OR c.name=:class_name)
                    GROUP BY f.id HAVING f.amount-COALESCE(SUM(p.amount),0)>0
                    ORDER BY s.surname,s.given_name,f.school_year,f.id""",
            ),
            "Paiements": (
                ["N° reçu", "Date", "Matricule", "Nom", "Prénom", "Classe", "Motif", "Montant (FCFA)"],
                """SELECT p.receipt_no,p.paid_at,s.matricule,s.surname,s.given_name,
                    COALESCE(c.name,''),p.reason,p.amount
                    FROM payments p JOIN students s ON s.id=p.student_id
                    LEFT JOIN classes c ON c.id=s.class_id
                    WHERE (:class_name IS NULL OR c.name=:class_name)
                    ORDER BY p.id DESC""",
            ),
            "Dépenses": (
                ["Date", "Catégorie", "Description", "Bénéficiaire", "Montant (FCFA)"],
                "SELECT spent_at,category,description,payee,amount FROM expenses ORDER BY id DESC",
            ),
            "Présences": (
                ["Date", "Matricule", "Nom", "Prénom", "Classe", "Statut", "Remarque"],
                """SELECT a.attendance_date,s.matricule,s.surname,s.given_name,COALESCE(c.name,''),a.status,a.note
                    FROM attendance a JOIN students s ON s.id=a.student_id
                    LEFT JOIN classes c ON c.id=s.class_id
                    WHERE (:class_name IS NULL OR c.name=:class_name)
                    ORDER BY a.attendance_date DESC,s.surname,s.given_name""",
            ),
            "Notes": (
                ["Année", "Période", "Matricule", "Nom", "Prénom", "Classe", "Matière", "Enseignant", "Note", "Sur", "Équivalent /10"],
                """SELECT g.school_year,g.period,s.matricule,s.surname,s.given_name,COALESCE(c.name,''),
                    sub.name,sub.teacher_name,g.score,g.max_score,ROUND(g.score*10.0/g.max_score,2)
                    FROM grades g JOIN students s ON s.id=g.student_id
                    JOIN subjects sub ON sub.id=g.subject_id
                    LEFT JOIN classes c ON c.id=s.class_id
                    WHERE (:class_name IS NULL OR c.name=:class_name)
                    ORDER BY g.school_year,g.period,s.surname,s.given_name,sub.name""",
            ),
        }
        if report not in queries:
            messagebox.showwarning("Rapport", "Choisissez un type de rapport.")
            return
        headers, query = queries[report]
        selected_class = self.report_class_filter.get()
        class_name = selected_class if selected_class and selected_class != "Toutes les classes" else None
        try:
            with connect() as db:
                params = {"class_name": class_name} if ":class_name" in query else {}
                rows = db.execute(query, params).fetchall()
        except sqlite3.Error as exc:
            messagebox.showerror("Rapport impossible", str(exc))
            return
        self.report_title = report
        self.report_headers = headers
        self.report_rows = [tuple(row) for row in rows]
        keys = [f"column_{i}" for i in range(len(headers))]
        self.report_tree.configure(columns=keys)
        self.report_tree.delete(*self.report_tree.get_children())
        for key, heading in zip(keys, headers):
            self.report_tree.heading(key, text=heading)
            self.report_tree.column(key, width=max(115, min(220, len(heading) * 11)), anchor="w", stretch=True)
        for row in self.report_rows:
            self.report_tree.insert("", "end", values=["" if value is None else value for value in row])
        scope = f" — classe {class_name}" if class_name else ""
        self.report_status.configure(text=f"{report}{scope} : {len(self.report_rows)} ligne(s).")

    def _report_export_path(self, extension: str) -> Path | None:
        if not self.report_headers:
            messagebox.showinfo("Aucun rapport", "Affichez d'abord un rapport avant de l'exporter.")
            return None
        report_dir = DATA_DIR / "rapports"
        report_dir.mkdir(parents=True, exist_ok=True)
        safe_title = self.report_title.lower().replace(" ", "_").replace("é", "e")
        target = filedialog.asksaveasfilename(
            title=f"Exporter le rapport en {extension.upper()}",
            initialdir=report_dir,
            initialfile=f"{safe_title}.{extension}",
            defaultextension=f".{extension}",
            filetypes=[(f"Fichier {extension.upper()}", f"*.{extension}")],
        )
        return Path(target) if target else None

    def export_report_csv(self) -> None:
        target = self._report_export_path("csv")
        if target is None:
            return
        try:
            with target.open("w", newline="", encoding="utf-8-sig") as stream:
                writer = csv.writer(stream, delimiter=";")
                writer.writerow(self.report_headers)
                writer.writerows(self.report_rows)
        except OSError as exc:
            messagebox.showerror("Export impossible", str(exc))
            return
        messagebox.showinfo("Export terminé", f"Le fichier CSV a été enregistré :\n{target}")

    def export_report_xlsx(self) -> None:
        target = self._report_export_path("xlsx")
        if target is None:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            messagebox.showerror("Composant absent", "Installez les dépendances avec : python -m pip install -r requirements.txt")
            return
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = self.report_title[:31]
        sheet.append(self.report_headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="16324F")
            cell.alignment = Alignment(horizontal="center")
        for row in self.report_rows:
            sheet.append(list(row))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            letter = column[0].column_letter
            max_width = max((len(str(cell.value or "")) for cell in column), default=10)
            sheet.column_dimensions[letter].width = min(max(max_width + 2, 12), 36)
        try:
            workbook.save(target)
        except OSError as exc:
            messagebox.showerror("Export impossible", str(exc))
            return
        messagebox.showinfo("Export terminé", f"Le fichier Excel a été enregistré :\n{target}")

    @staticmethod
    def copy_database(source_path: Path, target_path: Path) -> None:
        source = sqlite3.connect(source_path)
        destination = sqlite3.connect(target_path)
        try:
            source.backup(destination)
        finally:
            destination.close()
            source.close()

    @staticmethod
    def validate_database_file(path: Path) -> None:
        if not path.is_file():
            raise ValueError("Le fichier sélectionné n'existe pas.")
        source = sqlite3.connect(f"file:{path.resolve().as_posix()}?mode=ro", uri=True)
        try:
            result = source.execute("PRAGMA integrity_check").fetchone()
            if not result or result[0] != "ok":
                raise ValueError("La vérification d'intégrité de la base a échoué.")
            tables = {row[0] for row in source.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            required = {"classes", "students", "fees", "payments", "expenses", "audit_log"}
            missing = required - tables
            if missing:
                raise ValueError("Ce fichier ne semble pas être une sauvegarde du logiciel Sabiloul Houda.")
        finally:
            source.close()

    def create_backup(self) -> None:
        backup_dir = DATA_DIR / "sauvegardes"
        backup_dir.mkdir(parents=True, exist_ok=True)
        suggested = backup_dir / f"sabiloul_houda_{datetime.now():%Y%m%d_%H%M%S}.sqlite3"
        target = filedialog.asksaveasfilename(
            title="Enregistrer une sauvegarde",
            initialdir=backup_dir,
            initialfile=suggested.name,
            defaultextension=".sqlite3",
            filetypes=[("Sauvegarde SQLite", "*.sqlite3"), ("Tous les fichiers", "*.*")],
        )
        if not target:
            return
        if Path(target).resolve() == DB_PATH.resolve():
            messagebox.showwarning("Destination invalide", "Choisissez un autre nom que le fichier de base actuellement utilisé.")
            return
        try:
            self.copy_database(DB_PATH, Path(target))
            self.validate_database_file(Path(target))
        except (OSError, sqlite3.Error, ValueError) as exc:
            messagebox.showerror("Sauvegarde impossible", str(exc))
            return
        messagebox.showinfo("Sauvegarde créée", f"La copie de la base a été enregistrée ici :\n{target}")

    def restore_backup(self) -> None:
        source_name = filedialog.askopenfilename(
            title="Choisir une sauvegarde à restaurer",
            filetypes=[("Sauvegarde SQLite", "*.sqlite3 *.db"), ("Tous les fichiers", "*.*")],
        )
        if not source_name:
            return
        source_path = Path(source_name)
        if source_path.resolve() == DB_PATH.resolve():
            messagebox.showinfo("Déjà utilisée", "Ce fichier est déjà la base de données actuellement utilisée.")
            return
        try:
            self.validate_database_file(source_path)
        except (OSError, sqlite3.Error, ValueError) as exc:
            messagebox.showerror("Sauvegarde invalide", str(exc))
            return
        confirmed = messagebox.askyesno(
            "Confirmer la restauration",
            "La restauration remplacera les données actuellement utilisées.\n\n"
            "Une copie de sécurité de la base actuelle sera créée avant le remplacement. Continuer ?",
            icon="warning",
        )
        if not confirmed:
            return

        backup_dir = DATA_DIR / "sauvegardes"
        backup_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        safety_copy = backup_dir / f"avant_restauration_{stamp}.sqlite3"
        temporary_db = DATA_DIR / f"restauration_{stamp}.tmp"
        try:
            self.copy_database(DB_PATH, safety_copy)
            self.copy_database(source_path, temporary_db)
            self.validate_database_file(temporary_db)
            os.replace(temporary_db, DB_PATH)
            with connect() as db:
                db.execute(
                    "INSERT INTO audit_log(action,entity,details) VALUES('restauration','base de données',?)",
                    (source_path.name,),
                )
        except (OSError, sqlite3.Error, ValueError) as exc:
            if temporary_db.exists():
                temporary_db.unlink()
            messagebox.showerror(
                "Restauration impossible",
                f"{exc}\n\nLa copie de sécurité créée avant la restauration est :\n{safety_copy}",
            )
            return
        self.refresh_all()
        self.refresh_audit_log()
        messagebox.showinfo(
            "Restauration terminée",
            f"La base a été restaurée.\n\nCopie de sécurité de l'ancienne base :\n{safety_copy}",
        )

    def refresh_audit_log(self) -> None:
        with connect() as db:
            rows = db.execute(
                "SELECT happened_at,action,entity,entity_id,details FROM audit_log ORDER BY id DESC LIMIT 500"
            ).fetchall()
        self.clear(self.audit_tree)
        for row in rows:
            self.audit_tree.insert("", "end", values=tuple(row))

    @staticmethod
    def clear(tree: ttk.Treeview) -> None:
        tree.delete(*tree.get_children())

    def refresh_student_list(self, students=None) -> None:
        if students is None:
            with connect() as db:
                students = db.execute("""SELECT s.*, c.name AS class_name
                    FROM students s LEFT JOIN classes c ON c.id=s.class_id
                    ORDER BY s.surname COLLATE NOCASE,s.given_name COLLATE NOCASE""").fetchall()
        selected_class = self.student_filter_class.get()
        if selected_class and selected_class != "Toutes les classes":
            students = [row for row in students if row["class_name"] == selected_class]
        self.clear(self.student_tree)
        for row in students:
            self.student_tree.insert(
                "", "end", iid=str(row["id"]),
                values=(row["matricule"], row["surname"], row["given_name"], row["class_name"] or "", row["parent_phone"], row["status"]),
            )

    def refresh_all(self) -> None:
        with connect() as db:
            classes = db.execute("SELECT name FROM classes WHERE active=1 ORDER BY name").fetchall()
            self.student_class["values"] = [r["name"] for r in classes]
            if classes and not self.student_class.get():
                self.student_class.set(classes[0]["name"])
            students = db.execute("""SELECT s.*, c.name AS class_name FROM students s LEFT JOIN classes c ON c.id=s.class_id ORDER BY s.surname,s.given_name""").fetchall()
            filter_classes = db.execute("SELECT name FROM classes ORDER BY name COLLATE NOCASE").fetchall()
            filter_values = ["Toutes les classes", *(r["name"] for r in filter_classes)]
            self.student_filter_class["values"] = filter_values
            if self.student_filter_class.get() not in filter_values:
                self.student_filter_class.set("Toutes les classes")
            self.report_class_filter["values"] = filter_values
            if self.report_class_filter.get() not in filter_values:
                self.report_class_filter.set("Toutes les classes")
            self.update_report_class_filter()
            self.refresh_student_list(students)
            student_choices = [f"{r['id']} — {r['surname']} {r['given_name']} ({r['matricule']})" for r in students if r["status"] == "Actif"]
            self.fee_student["values"] = student_choices
            if self.fee_student.get() not in student_choices:
                self.fee_student.set(student_choices[0] if student_choices else "")
            class_rows = db.execute("SELECT c.id,c.name,c.active,COUNT(s.id) AS total FROM classes c LEFT JOIN students s ON s.class_id=c.id GROUP BY c.id ORDER BY c.name COLLATE NOCASE").fetchall()
            self.clear(self.class_tree)
            for r in class_rows:
                self.class_tree.insert("", "end", iid=str(r["id"]), values=(r["name"], r["total"], "Active" if r["active"] else "Archivée"))
            fees = db.execute("""SELECT f.id, f.school_year, f.label, f.amount,
                s.surname, s.given_name, s.matricule,
                COALESCE(SUM(p.amount),0) AS paid,
                f.amount-COALESCE(SUM(p.amount),0) AS remaining
                FROM fees f JOIN students s ON s.id=f.student_id
                LEFT JOIN payments p ON p.fee_id=f.id
                GROUP BY f.id ORDER BY s.surname,s.given_name,f.school_year,f.id""").fetchall()
            self.clear(self.fee_tree)
            due_choices = []
            for r in fees:
                student_name = f"{r['surname']} {r['given_name']} ({r['matricule']})"
                self.fee_tree.insert("", "end", iid=str(r["id"]), values=(student_name, r["school_year"], r["label"], f"{r['amount']:,}".replace(",", " "), f"{r['paid']:,}".replace(",", " "), f"{r['remaining']:,}".replace(",", " ")))
                if r["remaining"] > 0:
                    due_choices.append(f"{r['id']} — {student_name} — {r['label']} ({r['school_year']}) — reste {r['remaining']:,} FCFA".replace(",", " "))
            self.payment_fee["values"] = due_choices
            if self.payment_fee.get() not in due_choices:
                self.payment_fee.set(due_choices[0] if due_choices else "")
            self.update_payment_balance()
            payments = db.execute("SELECT p.*, s.surname, s.given_name FROM payments p JOIN students s ON s.id=p.student_id ORDER BY p.id DESC").fetchall()
            self.clear(self.payment_tree)
            for r in payments:
                self.payment_tree.insert("", "end", iid=str(r["id"]), values=(r["receipt_no"], r["paid_at"], f"{r['surname']} {r['given_name']}", r["reason"], f"{r['amount']:,}".replace(",", " ")))
            expenses = db.execute("SELECT * FROM expenses ORDER BY id DESC").fetchall()
            self.clear(self.expense_tree)
            for r in expenses:
                self.expense_tree.insert("", "end", iid=str(r["id"]), values=(r["spent_at"], r["category"], r["description"], r["payee"], f"{r['amount']:,}".replace(",", " ")))
            income = db.execute("SELECT COALESCE(SUM(amount),0) AS total FROM payments").fetchone()["total"]
            spent = db.execute("SELECT COALESCE(SUM(amount),0) AS total FROM expenses").fetchone()["total"]
            self.metrics["students"].configure(text=str(len(students)))
            self.metrics["income"].configure(text=f"{income:,}".replace(",", " "))
            self.metrics["expenses"].configure(text=f"{spent:,}".replace(",", " "))
            self.metrics["balance"].configure(text=f"{income-spent:,}".replace(",", " "))
        self.refresh_staff()
        with connect() as db:
            attendance_classes = [row["name"] for row in db.execute("SELECT name FROM classes WHERE active=1 ORDER BY name").fetchall()]
        self.attendance_class["values"] = attendance_classes
        if self.attendance_class.get() not in attendance_classes:
            self.attendance_class.set(attendance_classes[0] if attendance_classes else "")
        self.refresh_attendance_students()
        self.refresh_academics()
        bulletin_classes = [row["name"] for row in classes]
        self.bulletin_class["values"] = bulletin_classes
        if self.bulletin_class.get() not in bulletin_classes:
            self.bulletin_class.set(bulletin_classes[0] if bulletin_classes else "")
        self.refresh_bulletin_students()


if __name__ == "__main__":
    try:
        initialize_database()
        SchoolApp().mainloop()
    except (sqlite3.Error, OSError) as exc:
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror("Démarrage du logiciel", f"Impossible de démarrer le logiciel.\n\n{exc}")
        root.destroy()
